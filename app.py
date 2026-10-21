import os, io, re, calendar, sqlite3, threading, webbrowser
from collections import defaultdict
from datetime import date
from flask import Flask, request, redirect, render_template_string, g, send_file, flash, jsonify
from html.parser import HTMLParser
from urllib.parse import urlencode

DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), "hisab.db")
app = Flask(__name__)
app.secret_key = "hisab-local-ui"  # শুধু টোস্ট মেসেজ (flash) বহনের জন্য
app.jinja_env.filters["n"] = lambda x: f"{x:,.0f}"
app.jinja_env.filters["tc"] = lambda t: {"খরচ": "danger", "আয়": "success", "ধার নিলাম": "warning", "ধারে খরচ করলাম": "warning", "ধার পরিশোধ": "primary", "ধার দিলাম": "info", "ধার ফেরত পেলাম": "success"}.get(t, "secondary")
app.jinja_env.filters["z"] = lambda x: f"{x:,.0f}" if x else ""
def dmy(x):
    x = str(x or ""); return "%s-%s-%s" % (x[8:10], x[5:7], x[:4]) if len(x) >= 10 and x[4:5] == "-" else x
def my(x):
    x = str(x or ""); return "%s-%s" % (x[5:7], x[:4]) if len(x) == 7 and x[4:5] == "-" else x
def mname(x):
    x = str(x or ""); return "%s %s" % (calendar.month_name[int(x[5:7])], x[:4]) if re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", x) else x
def shift(m, d):
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", str(m or "")): m = date.today().strftime("%Y-%m")
    t = int(m[:4]) * 12 + int(m[5:7]) - 1 + d; return "%04d-%02d" % (t // 12, t % 12 + 1)
def curm(): return date.today().strftime("%Y-%m")
def getm():
    m = request.args.get("m", ""); return "" if m == "all" else m
def mq(m):
    a = request.args.to_dict(); a.pop("print", None)
    if m: a["m"] = m
    else: a.pop("m", None)
    return request.path + ("?" + urlencode(a) if a else "")
def ajax(): return request.headers.get("X-Requested-With") == "fetch"
def safe_next(nx, d="/"): return nx if nx and nx.startswith("/") and not nx.startswith("//") else d
def done(msg, nx, kind="success"):
    """সফল হলে টোস্ট মেসেজ রেখে (flash) নতুন পেজে পাঠায়; fetch হলে JSON ফেরত দেয়"""
    flash(msg, kind)
    return jsonify(ok=True, next=nx) if ajax() else redirect(nx)
def fmt(x): return f"{x:,.0f}"
def budget_warn(items):
    """items: [(date, type, category, amount)] — বাজেট ছুঁলে/ছাড়ালে টোস্টের বার্তার তালিকা"""
    add = defaultdict(float)
    for d_, t, c_, a in items:
        if t in NEED_CAT and c_: add[(str(d_)[:7], c_)] += a
    out, qs = [], ",".join("?" * len(NEED_CAT))
    for (mo, c_), a in add.items():
        lim = budgets_for(mo)[0].get(c_, 0)
        if not lim: continue
        sp = db().execute("SELECT COALESCE(SUM(amount),0) FROM entries WHERE category=? AND date LIKE ? AND type IN (%s)" % qs, (c_, mo + "%") + tuple(NEED_CAT)).fetchone()[0]
        tag = "" if mo == curm() else " (%s)" % my(mo)
        if sp > lim: out.append("%s খাতের বাজেট ছাড়িয়ে গেছে%s — %s / %s টাকা (%s বেশি)" % (c_, tag, fmt(sp), fmt(lim), fmt(sp - lim)))
        elif sp >= .8 * lim: out.append("%s খাতের বাজেটের %d%% খরচ হয়ে গেছে%s — %s / %s টাকা" % (c_, round(sp / lim * 100), tag, fmt(sp), fmt(lim)))
    return out
def budgets_for(mo):
    """প্রতি মাসের নিজস্ব বাজেট। ওই মাসে আলাদা করে সেট না থাকলে আগের সবচেয়ে কাছের সেট-করা মাসেরটা চলে।
    ফেরত: ({খাত: টাকা}, কোন মাস থেকে এসেছে বা None)"""
    src = db().execute("SELECT MAX(month) FROM mbudgets WHERE month<=?", (mo,)).fetchone()[0]
    if not src: return {}, None
    return {r[0]: r[1] for r in db().execute("SELECT category, amount FROM mbudgets WHERE month=? AND category!=''", (src,))}, src
def month_budget(mm):
    """বাছা মাসের খাতভিত্তিক খরচ বনাম বাজেট"""
    spent = defaultdict(float)
    for e in entries(mm):
        if e["type"] in NEED_CAT: spent[e["category"]] += e["amount"]
    bd, src = budgets_for(mm); rows = []
    for c_, lim in bd.items():
        sp = spent.get(c_, 0); p = sp / lim * 100 if lim else 0
        rows.append(dict(c=c_, l=lim, s=sp, p=p, st="over" if p > 100 else "warn" if p >= 80 else "ok"))
    rows.sort(key=lambda x: -x["p"])
    return rows, spent, src
def goal_get(mo):
    """প্রতি মাসের নিজস্ব সঞ্চয় লক্ষ্য; ওই মাসে সেট না থাকলে আগের সবচেয়ে কাছের সেট-করা মাসেরটা চলে। ফেরত: (টাকা, উৎস মাস)"""
    src = db().execute("SELECT MAX(month) FROM mgoals WHERE month<=?", (mo,)).fetchone()[0]
    if not src: return 0, None
    return db().execute("SELECT amount FROM mgoals WHERE month=?", (src,)).fetchone()[0] or 0, src
app.jinja_env.filters["dmy"] = dmy; app.jinja_env.filters["my"] = my; app.jinja_env.filters["mname"] = mname
app.jinja_env.globals.update(shift=shift, curm=curm, mq=mq)

TYPES = ["খরচ", "আয়", "ধার নিলাম", "ধারে খরচ করলাম", "ধার পরিশোধ", "ধার দিলাম", "ধার ফেরত পেলাম"]
NEED_CAT = {"খরচ", "ধারে খরচ করলাম"}
NEED_NAME = set(TYPES) - {"খরচ", "আয়"}
MEDIA = ["Cash", "Bank", "bKash"]
CASH = {"খরচ": -1, "আয়": 1, "ধার নিলাম": 1, "ধারে খরচ করলাম": 0, "ধার পরিশোধ": -1, "ধার দিলাম": -1, "ধার ফেরত পেলাম": 1}

def init():
    c = sqlite3.connect(DB)
    c.execute("CREATE TABLE IF NOT EXISTS entries(id INTEGER PRIMARY KEY, date TEXT, type TEXT, medium TEXT, category TEXT, description TEXT, name TEXT, amount REAL)")
    c.execute("CREATE TABLE IF NOT EXISTS opening(medium TEXT PRIMARY KEY, amount REAL)")
    c.execute("CREATE TABLE IF NOT EXISTS renames(old TEXT PRIMARY KEY, new TEXT)")
    c.execute("CREATE TABLE IF NOT EXISTS budgets(category TEXT PRIMARY KEY, amount REAL)")
    c.execute("CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value REAL)")
    c.execute("CREATE TABLE IF NOT EXISTS mgoals(month TEXT PRIMARY KEY, amount REAL)")
    if not c.execute("SELECT 1 FROM mgoals LIMIT 1").fetchone():  # আগের একক লক্ষ্য '0000-00' নামে বেস হিসেবে থাকবে
        c.execute("INSERT OR IGNORE INTO mgoals SELECT '0000-00', value FROM settings WHERE key='goal'")
    c.execute("CREATE TABLE IF NOT EXISTS mbudgets(month TEXT, category TEXT, amount REAL, PRIMARY KEY(month, category))")
    if not c.execute("SELECT 1 FROM mbudgets LIMIT 1").fetchone():  # আগের সাধারণ বাজেট '0000-00' নামে বেস হিসেবে থাকবে
        c.execute("INSERT OR IGNORE INTO mbudgets SELECT '0000-00', category, amount FROM budgets")
    for m in MEDIA:
        c.execute("INSERT OR IGNORE INTO opening VALUES(?,0)", (m,))
    c.commit(); c.close()
init()

def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB); g.db.row_factory = sqlite3.Row
    return g.db

@app.teardown_appcontext
def close(_):
    d = g.pop("db", None)
    if d: d.close()

def pair(e):
    """প্রতিটি এন্ট্রি থেকে অটো জার্নাল: (Debit account, Credit account)"""
    t, m, cat, nm = e["type"], e["medium"], e["category"] or "খাত নেই", e["name"] or "নাম নেই"
    x, i, p, r = "খরচ: " + cat, "আয়: " + (e["category"] or "উৎস নেই"), "Payable: " + nm, "Receivable: " + nm
    return {"খরচ": (x, m), "আয়": (m, i), "ধার নিলাম": (m, p), "ধারে খরচ করলাম": (x, p),
            "ধার পরিশোধ": (p, m), "ধার দিলাম": (r, m), "ধার ফেরত পেলাম": (m, r)}[t]

def kind(a):
    if a in MEDIA or a.startswith("Receivable"): return "A"
    if a.startswith("Payable"): return "L"
    if a.startswith("খরচ"): return "X"
    if a.startswith("আয়"): return "I"
    return "E"
KN = {"A": "সম্পদ (Asset)", "L": "দায় (Payable)", "E": "মূলধন", "I": "আয়", "X": "খরচ"}

def entries(m=None):
    q, a = "SELECT * FROM entries", ()
    if m: q, a = q + " WHERE date LIKE ?", (m + "%",)
    return db().execute(q + " ORDER BY date, id", a).fetchall()

def months():
    return [r[0] for r in db().execute("SELECT DISTINCT substr(date,1,7) FROM entries ORDER BY 1 DESC")]

def books(upto=None):
    ops = {r["medium"]: r["amount"] for r in db().execute("SELECT * FROM opening")}
    bal = defaultdict(lambda: [0.0, 0.0])
    for m, a in ops.items():
        bal[m][0] += a; bal["Opening Capital"][1] += a
    rows = []
    for e in entries():
        if upto and e["date"][:7] > upto: continue
        d, c = pair(e); rows.append((e, d, c))
        bal[d][0] += e["amount"]; bal[c][1] += e["amount"]
    return ops, rows, bal

BASE = """<!doctype html><html lang="bn" data-theme="green"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{{title}} — হিসাব</title>
<script>try{var d0=document.documentElement,t=localStorage.getItem('th')||(matchMedia('(prefers-color-scheme:dark)').matches?'dark':'light');d0.setAttribute('data-bs-theme',t);d0.setAttribute('data-theme',localStorage.getItem('tm')||'green')}catch(e){}</script>
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Hind+Siliguri:wght@400;500;600;700&family=Inter:wght@400;500;600;700&family=Poppins:wght@500;600;700&display=swap" rel="stylesheet">
<link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/css/bootstrap.min.css" rel="stylesheet">
<link href="https://cdn.jsdelivr.net/npm/flatpickr@4.6.13/dist/flatpickr.min.css" rel="stylesheet">
<link href="https://cdn.jsdelivr.net/npm/bootstrap-icons@1.11.3/font/bootstrap-icons.min.css" rel="stylesheet">
<style>
:root{--b1:#2c572a;--b2:#8cc542;--b1rgb:44,87,42;--ac-rgb:var(--b1rgb);--ac:var(--b1);--bg:#f2f7ee;--card:#fff;--text:#1b2a1b;--muted:#5d6e5d;--line:#e0e9d8;--brand:var(--ac);--shadow:0 4px 18px rgba(30,60,30,.07);--bs-primary:var(--b1);--bs-primary-rgb:var(--b1rgb);--bs-link-color:var(--ac);--bs-link-hover-color:var(--ac)}
[data-bs-theme=dark]{--ac:var(--b2);--ac-rgb:140,197,66;--bg:#0c170f;--card:#14231a;--text:#e3eddc;--muted:#9bb09a;--line:#244029;--shadow:0 4px 18px rgba(0,0,0,.4);--bs-link-color:var(--ac);--bs-link-hover-color:var(--ac)}
body{font-family:'Inter','Hind Siliguri','Noto Sans Bengali',system-ui,sans-serif;background:var(--bg);color:var(--text)}
.side{position:fixed;inset:0 auto 0 0;width:240px;background:var(--card);border-right:1px solid var(--line);padding:20px 14px;display:flex;flex-direction:column;gap:4px;z-index:10;overflow-y:auto}
.brand{display:flex;align-items:center;gap:10px;font-size:1.35rem;font-weight:700;padding:0 10px 18px;color:var(--text)}
.brand i{background:linear-gradient(135deg,var(--b1),var(--b2));color:#fff;border-radius:12px;padding:6px 9px;font-size:1.1rem}
.side a,.side button{display:flex;align-items:center;gap:12px;padding:10px 14px;border-radius:12px;color:var(--muted);text-decoration:none;font-weight:500;border:0;background:none;width:100%;text-align:left;white-space:nowrap}
.side a:hover,.side button:hover{background:var(--bg);color:var(--text)}
.side a.on{background:linear-gradient(135deg,var(--b1),var(--b2));color:#fff;box-shadow:0 6px 16px rgba(var(--b1rgb),.35)}
.side .tg{margin-top:auto}
.main{margin-left:240px;padding:28px 32px 60px}
.ttl{font-weight:700;font-size:1.55rem;margin-bottom:18px}
.stat{position:relative;border-radius:18px;padding:16px;color:#fff;overflow:hidden;box-shadow:var(--shadow)}
.stat i{position:absolute;right:12px;top:10px;font-size:2rem;opacity:.28}
.stat .lbl{font-size:.82rem;opacity:.92}.stat .val{font-size:1.6rem;font-weight:700;letter-spacing:-.5px}
.g1{background:linear-gradient(135deg,var(--b1),var(--b2))}
.g2{background:linear-gradient(135deg,#1c6b4d,#3fa66b);background:linear-gradient(135deg,color-mix(in srgb,var(--b1) 65%,#0d9488),color-mix(in srgb,var(--b2) 65%,#2dd4bf))}
.g3{background:linear-gradient(135deg,#4f8f2a,#8cc542);background:linear-gradient(135deg,color-mix(in srgb,var(--b1) 78%,var(--b2)),var(--b2))}
.g4{background:linear-gradient(135deg,#b45309,#f59e0b)}
.g5{background:linear-gradient(135deg,#2f7d32,#8cc542);background:linear-gradient(135deg,color-mix(in srgb,var(--b1) 88%,#000),color-mix(in srgb,var(--b2) 92%,#000))}
.g6{background:linear-gradient(135deg,#1d3b1d,#2c572a);background:linear-gradient(135deg,color-mix(in srgb,var(--b1) 55%,#000),var(--b1))}
.card{background:var(--card);border:1px solid var(--line);border-radius:18px;box-shadow:var(--shadow)}
.card-header{background:transparent;border-bottom:1px solid var(--line);font-weight:600;padding:14px 18px}
.card-body{padding:18px}
.table{--bs-table-bg:var(--card);--bs-table-color:var(--text);--bs-table-striped-bg:rgba(100,116,139,.06);--bs-table-hover-bg:rgba(var(--ac-rgb),.08);border-collapse:separate;border-spacing:0;border-radius:14px;overflow:hidden;box-shadow:var(--shadow);border:1px solid var(--line)}
.table thead th{background:rgba(var(--ac-rgb),.09);color:var(--muted);font-weight:600;font-size:.8rem;border-bottom:1px solid var(--line)}
.table td,.table th{padding:.6rem .85rem;border-color:var(--line)}
.form-control,.form-select{border-radius:12px;border-color:var(--line);background-color:var(--card);color:var(--text);padding:.6rem .9rem}
.form-control:focus,.form-select:focus{border-color:var(--brand);box-shadow:0 0 0 .2rem rgba(var(--ac-rgb),.2)}
.btn-primary{background:linear-gradient(135deg,var(--b1),var(--b2));border:0;border-radius:12px;font-weight:600;padding:.6rem 1rem}
.btn-primary:hover{filter:brightness(1.08)}
.badge{border-radius:20px;font-weight:500;padding:.4em .75em}
.alert{border-radius:14px}
#qa .modal-header{padding:.6rem 1rem}#qa .modal-title{font-size:1rem}#qa .modal-body{padding:.9rem 1rem}#qa .form-control,#qa .form-select{padding:.35rem .65rem;font-size:.88rem;border-radius:10px}#qa .table td,#qa .table th{padding:.3rem .45rem;font-size:.85rem}#fB select.form-select-sm{min-width:110px;padding-right:1.9rem;background-position:right .55rem center}#fB td{white-space:nowrap}
@media(max-width:991px){.side{position:static;width:auto;flex-direction:row;overflow-x:auto;padding:10px;border-right:0;border-bottom:1px solid var(--line)}
.side a,.side button{width:auto;padding:8px 12px}.brand{display:none}.side .tg{margin-top:0}.main{margin-left:0;padding:16px}.ttl{font-size:1.3rem}}
.sec{font-size:.68rem;font-weight:700;letter-spacing:.08em;color:var(--muted);padding:14px 14px 4px;text-transform:uppercase}
.dots{display:flex;flex-wrap:wrap;gap:9px;padding:2px 14px 6px}
.side .dot{width:24px;height:24px;min-width:24px;padding:0;border-radius:50%;border:0;cursor:pointer;box-shadow:0 0 0 2px var(--card),0 0 0 3px var(--line);transition:transform .15s,box-shadow .15s}
.side .dot:hover{transform:scale(1.15)}.side .dot.on{box-shadow:0 0 0 2px var(--card),0 0 0 4px var(--ac)}.side .dot:focus-visible{outline:2px solid var(--ac);outline-offset:3px}
.fab{position:fixed;right:24px;bottom:24px;width:56px;height:56px;border-radius:50%;background:linear-gradient(135deg,var(--b1),var(--b2));color:#fff;border:0;font-size:1.5rem;box-shadow:0 10px 24px rgba(var(--b1rgb),.45);z-index:20}
.fab:hover{transform:scale(1.08)}.fab{transition:transform .15s}
.main>*{animation:fi .35s ease both}@keyframes fi{from{opacity:0;transform:translateY(6px)}to{opacity:1;transform:none}}
.stat{transition:transform .15s}.stat:hover{transform:translateY(-3px)}
.prh{display:none;margin-bottom:8px;font-size:.9rem;color:#475569}
@media(max-width:991px){.sec{display:none}.dots{padding:0 6px}.side .tg{display:flex;align-items:center}}
@media print{*{-webkit-print-color-adjust:exact;print-color-adjust:exact}.side,.fab,.tb,.modal,.btn,.noprint,.toasts{display:none!important}.main{margin:0!important;padding:0!important}body{background:#fff!important;color:#000}
.prh{display:block}.sw{max-height:none!important;overflow:visible!important;box-shadow:none}.table{box-shadow:none}.table.sh td,.table.sh th{font-size:6.5pt;padding:1px 3px}
.table.sh td:first-child,.table.sh th:first-child,.table.sh thead th{position:static}@page{size:A4 portrait;margin:8mm}}
.mnav{display:inline-flex;align-items:center;gap:4px;background:var(--card);border:1px solid var(--line);border-radius:40px;padding:5px 7px;box-shadow:var(--shadow)}
.mb-btn{width:38px;height:38px;border-radius:50%;display:grid;place-items:center;color:var(--text);text-decoration:none;font-size:1.05rem;transition:background .15s,transform .15s}
.mb-btn:hover{background:var(--bg);transform:scale(1.1)}.mb-label{min-width:180px;text-align:center;font-weight:700;font-size:1.1rem;padding:0 6px}
.mb-chip{font-size:.8rem;padding:5px 14px;border-radius:20px;color:var(--muted);text-decoration:none;border:1px solid var(--line);background:var(--card)}
.mb-chip:hover{color:var(--text)}.mb-chip.on{background:linear-gradient(135deg,var(--b1),var(--b2));color:#fff;border-color:transparent}
@media print{.mnav,.mb-chip{display:none!important}}
.stat{text-shadow:0 1px 2px rgba(0,0,0,.25)}
.btn-outline-primary{--bs-btn-color:var(--ac);--bs-btn-border-color:var(--ac);--bs-btn-hover-bg:var(--b1);--bs-btn-hover-border-color:var(--b1);--bs-btn-hover-color:#fff;--bs-btn-active-bg:var(--b1);--bs-btn-active-border-color:var(--b1)}
.btn-success,.btn-outline-success{--bs-btn-color:var(--ac);--bs-btn-border-color:var(--ac);--bs-btn-hover-bg:var(--b1);--bs-btn-hover-border-color:var(--b1);--bs-btn-hover-color:#fff;--bs-btn-active-bg:var(--b1);--bs-btn-active-border-color:var(--b1)}
.btn-success{--bs-btn-bg:var(--b1);--bs-btn-color:#fff;--bs-btn-border-color:var(--b1)}
.fixbar{position:fixed;top:0;left:240px;right:0;z-index:30;background:var(--bg);padding:14px 32px 6px;box-shadow:0 10px 16px -14px rgba(0,0,0,.35)}
.fixbar .ttl{font-size:1.35rem;margin:0}
.fixbar .mnav{padding:3px 5px}.fixbar .mb-btn{width:32px;height:32px}.fixbar .mb-label{min-width:150px;font-size:1rem}.fixbar .mb-chip{padding:3px 12px}
.fixbar .stat{padding:9px 13px;border-radius:14px}.fixbar .stat .val{font-size:1.3rem}.fixbar .stat i{font-size:1.5rem;top:8px}
.fixbar .mb-3,.fixbar .mb-4{margin-bottom:.55rem!important}.fixbar .row.g-3{--bs-gutter-y:.6rem;--bs-gutter-x:.75rem}
@media(max-width:991px){.fixbar{position:static;margin:0 0 12px;padding:0;box-shadow:none}}
@media print{.fixbar{position:static;box-shadow:none;margin:0;padding:0}}

/* ---- v16: typography ---- */
.ttl,.brand,.stat .val,.mb-label,.card-header,.modal-title{font-family:'Poppins','Inter','Hind Siliguri','Noto Sans Bengali',system-ui,sans-serif}
.ttl{letter-spacing:-.01em}
.table td,.table th,.stat .val,.form-control,.mb-label{font-variant-numeric:tabular-nums}
/* ---- v16: colour themes (light + dark each) ---- */
[data-theme=midnight]{--b1:#1e3a8a;--b2:#3b82f6;--b1rgb:30,58,138;--bg:#eef2fb;--line:#d9e1f2;--text:#0f1a33;--muted:#5b6a8a}
[data-theme=midnight][data-bs-theme=dark]{--bg:#060b1c;--card:#0d1630;--line:#1c2a52;--text:#dde6fa;--muted:#8a9cc4;--ac-rgb:96,165,250}
[data-theme=emerald]{--b1:#047857;--b2:#34d399;--b1rgb:4,120,87;--bg:#eef8f4;--line:#d3eadf;--text:#0f2a22;--muted:#52716a}
[data-theme=emerald][data-bs-theme=dark]{--bg:#061512;--card:#0d241e;--line:#1a3d34;--text:#dcf1ea;--muted:#8fb5aa;--ac-rgb:52,211,153}
[data-theme=sunset]{--b1:#c2410c;--b2:#f97316;--b1rgb:194,65,12;--bg:#fdf4ee;--line:#f3dfd0;--text:#2e1a10;--muted:#85624d}
[data-theme=sunset][data-bs-theme=dark]{--bg:#180d08;--card:#25140c;--line:#4a2a18;--text:#f6e4d8;--muted:#c0a08c;--ac-rgb:251,146,60}
[data-theme=sunset] .g4{background:linear-gradient(135deg,#9f1239,#e11d48)}
[data-theme=violet]{--b1:#5b21b6;--b2:#8b5cf6;--b1rgb:91,33,182;--bg:#f5f1fc;--line:#e5dcf5;--text:#1e1433;--muted:#6b5c8a}
[data-theme=violet][data-bs-theme=dark]{--bg:#0e0919;--card:#1a1230;--line:#32214f;--text:#e9e1fa;--muted:#a99bc9;--ac-rgb:167,139,250}
[data-theme=rose]{--b1:#be123c;--b2:#f43f5e;--b1rgb:190,18,60;--bg:#fdf1f3;--line:#f5d9de;--text:#2d1117;--muted:#86545e}
[data-theme=rose][data-bs-theme=dark]{--bg:#160a0e;--card:#241119;--line:#46202c;--text:#f8e3e8;--muted:#c4959f;--ac-rgb:251,113,133}
[data-theme=rose] .g4{background:linear-gradient(135deg,#b45309,#f59e0b)}
html.tmx *,html.tmx *::before{transition:background-color .35s,border-color .35s,color .25s!important}
/* ---- v16: motion (answers an action: open, switch, confirm) ---- */
.main>.fixbar{animation:none}
.main>*:nth-child(n+3){animation-delay:.05s}.main>*:nth-child(n+5){animation-delay:.1s}
.modal.fade .modal-dialog{transform:translateY(16px) scale(.96);opacity:0;transition:transform .3s cubic-bezier(.2,.85,.25,1),opacity .22s ease}
.modal.show .modal-dialog{transform:none;opacity:1}
.modal-backdrop.show{opacity:.5;backdrop-filter:blur(3px)}
.pane-in{animation:pin .26s cubic-bezier(.2,.8,.2,1)}@keyframes pin{from{opacity:0;transform:translateX(10px)}to{opacity:1;transform:none}}
#rows tr{animation:rin .24s ease}@keyframes rin{from{opacity:0;transform:translateY(-4px)}to{opacity:1;transform:none}}
.btn{transition:transform .12s,filter .15s,background-color .15s,border-color .15s,color .15s,box-shadow .15s}.btn:active{transform:scale(.97)}
.cf-ic{width:56px;height:56px;border-radius:50%;margin:0 auto 12px;display:grid;place-items:center;font-size:1.5rem;background:rgba(225,29,72,.12);color:#e11d48}
/* ---- v16: toast notifications ---- */
.toasts{position:fixed;left:264px;bottom:24px;z-index:1100;display:flex;flex-direction:column;gap:10px;width:min(390px,calc(100vw - 32px));pointer-events:none}
.tst{--tc:var(--ac-rgb);pointer-events:auto;position:relative;display:flex;align-items:center;gap:12px;padding:12px 12px 15px 14px;background:var(--card);color:var(--text);border:1px solid var(--line);border-radius:14px;box-shadow:0 14px 34px rgba(0,0,0,.2);overflow:hidden;animation:tin .34s cubic-bezier(.2,.9,.25,1.1) both}
.tst-error{--tc:217,119,6}.tst-delete{--tc:225,29,72}.tst-info{--tc:14,165,233}
.tst>i{flex:none;width:34px;height:34px;border-radius:50%;display:grid;place-items:center;background:rgba(var(--tc),.14);color:rgb(var(--tc))}
.tst .tm{flex:1;font-weight:500;font-size:.92rem;line-height:1.45}
.tst .tcl{flex:none;border:0;background:none;color:var(--muted);padding:5px 6px;border-radius:8px;line-height:1}.tst .tcl:hover{background:var(--bg);color:var(--text)}
.tst .tp{position:absolute;left:0;bottom:0;height:3px;width:100%;background:rgb(var(--tc));transform-origin:left;animation:tp linear forwards;animation-duration:var(--d,3.6s)}
.tst:hover .tp{animation-play-state:paused}
.tst.out{animation:tout .2s ease forwards}
@keyframes tin{from{opacity:0;transform:translateX(-20px) scale(.97)}to{opacity:1;transform:none}}
@keyframes tout{to{opacity:0;transform:translateX(-14px) scale(.97)}}
@keyframes tp{from{transform:scaleX(1)}to{transform:scaleX(0)}}
@media(max-width:991px){.toasts{left:12px;right:92px;width:auto;bottom:20px}#tname{display:none}}
@media(prefers-reduced-motion:reduce){*,*::before,*::after{animation-duration:.01ms!important;animation-delay:0s!important;transition-duration:.01ms!important}.tst .tp{animation:tp var(--d,3.6s) linear forwards!important}}
/* ---- v17: budget + savings progress ---- */
.tst-warn{--tc:234,88,12}
.pgt{position:relative;height:16px;border-radius:99px;background:rgba(var(--ac-rgb),.13);overflow:hidden}
.pgt.sm{height:8px;margin:6px 0 2px}
.pg{height:100%;width:0;border-radius:99px;transition:width 1s cubic-bezier(.2,.8,.2,1)}
.pg.ok{background:linear-gradient(90deg,var(--b1),var(--b2))}
.pg.warn{background:linear-gradient(90deg,#d97706,#f59e0b)}
.pg.over{background:linear-gradient(90deg,#be123c,#f43f5e)}
.pg.done{box-shadow:0 0 0 2px rgba(var(--ac-rgb),.25) inset}
.sv-big{font-family:'Poppins','Inter','Hind Siliguri',sans-serif;font-size:2rem;font-weight:700;letter-spacing:-.02em;line-height:1}
.bd-row{margin-bottom:14px}.bd-row:last-child{margin-bottom:0}
.bd-ok{color:var(--muted)}.bd-warn{color:#d97706}.bd-over{color:#e11d48}
.sv-hist{display:flex;gap:10px;align-items:flex-end}
.sv-col{flex:1;text-align:center;font-size:.72rem;color:var(--muted)}
.sv-col.cur{color:var(--text);font-weight:600}
.sv-bar{height:56px;display:flex;align-items:flex-end;border-top:1px dashed rgba(var(--ac-rgb),.55);margin-bottom:4px}
.sv-bar i{display:block;width:100%;border-radius:6px 6px 2px 2px;transform-origin:bottom;animation:grow .7s cubic-bezier(.2,.8,.2,1) both}
.sv-bar i.full{background:linear-gradient(180deg,var(--b2),var(--b1))}.sv-bar i.part{background:rgba(var(--ac-rgb),.35)}
@keyframes grow{from{transform:scaleY(0)}to{transform:scaleY(1)}}
</style></head><body>
<nav class="side"><div class="brand"><i class="bi bi-wallet2"></i><span class="bn">হিসাব</span></div>
{% for g,items in [('Main',[('/','speedometer2','Dashboard'),('/entries','list-ul','Entries'),('/sheet','table','Monthly Sheet'),('/heads','tags','Edit Heads'),('/budget','piggy-bank','Budget & Goal')]),('Books',[('/journal','journal-text','Journal'),('/ledger','book','Ledger'),('/trial','calculator','Trial Balance')]),('Reports',[('/monthly','file-earmark-bar-graph','Monthly Report'),('/people','people','Loan Summary'),('/balance','wallet2','Balance Summary'),('/downloads','download','Download Center')])] %}
<div class="sec">{{g}}</div>{% for u,ic,l in items %}<a href="{{u}}" class="{{'on' if request.path==u}}"><i class="bi bi-{{ic}}"></i>{{l}}</a>{% endfor %}{% endfor %}
<div class="tg"><div class="sec noprint">Theme</div><div class="dots" id="dots" role="radiogroup" aria-label="Color theme"></div><div class="small text-muted px-3 pb-2 noprint" id="tname"></div><button onclick="tgl()"><i class="bi bi-circle-half"></i>Dark / Light</button><div class="small text-muted px-3 pt-1 noprint">Version: Report v15 ({{fname}})</div></div></nav>
<main class="main"><div class="prh">হিসাব — {{title}} — {{today|dmy}}</div>
<div class="fixbar"><div class="tb d-flex flex-wrap align-items-center gap-2 mb-2"><div class="ttl m-0 me-auto">{{title}}</div>
<input id="q" class="form-control form-control-sm w-auto" placeholder="Search this page…">
<button class="btn btn-sm btn-outline-secondary" onclick="pal()" title="Ctrl + K"><i class="bi bi-command"></i> Menu / Search</button>
<a class="btn btn-sm btn-outline-success" href="/export?u={{request.full_path|urlencode}}"><i class="bi bi-file-earmark-excel"></i> Excel (this page)</a>
<a class="btn btn-sm btn-success" href="/export_all?m={{request.args.get('m','')}}" title="Monthly Sheet, Monthly Report, Loan Summary ও Balance Summary — ৪টি শিট এক ফাইলে"><i class="bi bi-file-earmark-spreadsheet"></i> All Excel (1 click)</a>
<button class="btn btn-sm btn-outline-danger" onclick="print()"><i class="bi bi-file-earmark-pdf"></i> PDF</button>
<a class="btn btn-sm btn-outline-primary" href="/backup" onclick="toast('ব্যাকআপ ডাউনলোড শুরু হয়েছে','info')"><i class="bi bi-cloud-arrow-down"></i> Backup</a></div>{{fix|safe}}</div>
{{body|safe}}</main>
<button class="fab" data-bs-toggle="modal" data-bs-target="#qa" title="New Entry"><i class="bi bi-plus-lg"></i></button>
<div class="modal fade" id="qa"><div class="modal-dialog modal-dialog-centered modal-dialog-scrollable modal-lg"><div class="modal-content" style="border-radius:16px;background:var(--card)">
<div class="modal-header"><h5 class="modal-title" id="qtitle">New Entry</h5><button type="button" class="btn btn-sm btn-outline-primary mx-auto" id="tabB" onclick="qtab('b')"><i class="bi bi-table"></i> Bulk Entry</button><button type="button" class="btn btn-sm btn-outline-primary mx-auto d-none" id="tabS" onclick="qtab('s')"><i class="bi bi-pencil-square"></i> Single Entry</button><button class="btn-close" data-bs-dismiss="modal"></button></div>
<form id="fS" method="post" action="/add" data-ajax class="modal-body row g-2"><input type="hidden" name="next" value="{{request.full_path}}">
<div class="col-6"><input type="text" data-date name="date" value="{{today}}" class="form-control" required></div>
<div class="col-6"><select name="type" class="form-select" onchange="syncS()">{% for t in types %}<option>{{t}}</option>{% endfor %}</select></div>
<div class="col-6"><select name="medium" class="form-select">{% for m in media %}<option>{{m}}</option>{% endfor %}</select></div>
<div class="col-6" id="sCat"><input name="category" list="qc" class="form-control" placeholder="Category / Income source"></div>
<div class="col-6 d-none" id="sName"><input name="name" list="qn" class="form-control" placeholder="Name (for loans / dues)"></div>
<div class="col-6"><input name="description" class="form-control" placeholder="Description"></div>
<div class="col-6"><input type="number" step="any" name="amount" class="form-control" placeholder="Amount" required></div>
<div class="col-6"><button class="btn btn-primary w-100">Save</button></div>
</form>
<form id="fB" method="post" action="/add_bulk" data-ajax class="modal-body d-none"><input type="hidden" name="next" value="{{request.full_path}}">
<div class="mb-2"><label class="small text-muted">Date (for all rows)</label><input type="text" data-date name="date" value="{{today}}" class="form-control form-control-sm" style="max-width:220px"></div>
<div class="table-responsive"><table class="table table-sm align-middle mb-2"><thead><tr><th>Type<th>Medium<th>Category / Source<th>Name<th>Description<th class="text-end">Amount<th></tr></thead><tbody id="rows"></tbody></table></div>
<button type="button" class="btn btn-sm btn-outline-primary" onclick="addRow()"><i class="bi bi-plus-lg"></i> Add row</button>
<template id="rowtpl"><tr><td><select name="type[]" class="form-select form-select-sm" onchange="syncRow(this)">{% for t in types %}<option>{{t}}</option>{% endfor %}</select><td><select name="medium[]" class="form-select form-select-sm">{% for m in media %}<option>{{m}}</option>{% endfor %}</select><td><input name="category[]" list="qc" class="form-control form-control-sm"><td><input name="name[]" list="qn" class="form-control form-control-sm"><td><input name="description[]" class="form-control form-control-sm"><td><input type="number" step="any" name="amount[]" class="form-control form-control-sm text-end"><td><button type="button" class="btn btn-sm btn-outline-danger" onclick="this.closest('tr').remove()" title="Delete this row"><i class="bi bi-x-lg"></i></button></template>
<div class="modal-footer border-0 px-0 pb-0 mt-2"><button type="button" class="btn btn-outline-secondary" data-bs-dismiss="modal">Cancel</button><button class="btn btn-primary">Save all</button></div>
</form>
<datalist id="qc">{% for c in qc %}<option value="{{c}}">{% endfor %}</datalist><datalist id="qn">{% for c in qn %}<option value="{{c}}">{% endfor %}</datalist></div></div></div>
<div class="modal fade" id="pal"><div class="modal-dialog modal-dialog-centered"><div class="modal-content" style="border-radius:18px;background:var(--card)">
<div class="p-3"><input id="pi" class="form-control" placeholder="Where to go, or what to do? Type here…" autocomplete="off"></div><div id="pl" class="list-group list-group-flush pb-2"></div></div></div></div>
<div class="toasts" id="toasts" aria-live="polite"></div>
<div class="modal fade" id="cf" tabindex="-1"><div class="modal-dialog modal-dialog-centered modal-sm"><div class="modal-content" style="border-radius:18px;background:var(--card)"><div class="modal-body text-center p-4">
<div class="cf-ic"><i class="bi bi-trash3"></i></div><div class="fw-bold fs-5 mb-1" id="cft">মুছবেন?</div><div class="fw-semibold small mb-1" id="cfs"></div><div class="small text-muted mb-3">মুছে ফেললে ফিরিয়ে আনা যাবে না। সব রিপোর্ট নিজে থেকেই ঠিক হয়ে যাবে।</div>
<div class="d-flex gap-2 justify-content-center"><button type="button" class="btn btn-outline-secondary" data-bs-dismiss="modal">বাতিল</button><button type="button" class="btn btn-danger" id="cfy"><i class="bi bi-trash3"></i> হ্যাঁ, মুছুন</button></div></div></div></div></div>
<script>var FL={{ get_flashed_messages(with_categories=true)|tojson }};</script>
<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/js/bootstrap.bundle.min.js"></script>
<script src="https://cdn.jsdelivr.net/npm/flatpickr@4.6.13/dist/flatpickr.min.js"></script>
<script>
var THEMES=[{k:'green',n:'ক্লাসিক গ্রিন',c:['#2c572a','#8cc542']},{k:'midnight',n:'মিডনাইট ব্লু',c:['#1e3a8a','#3b82f6']},{k:'emerald',n:'এমারেল্ড গ্রিন',c:['#047857','#34d399']},{k:'sunset',n:'সানসেট অরেঞ্জ',c:['#c2410c','#f97316']},{k:'violet',n:'রয়্যাল ভায়োলেট',c:['#5b21b6','#8b5cf6']},{k:'rose',n:'রোজ',c:['#be123c','#f43f5e']}];
function smooth(fn){var h=document.documentElement;h.classList.add('tmx');fn(h);clearTimeout(window._tx);window._tx=setTimeout(function(){h.classList.remove('tmx')},450);if(window.reTheme)reTheme()}
function tgl(){smooth(function(h){var t=h.getAttribute('data-bs-theme')=='dark'?'light':'dark';h.setAttribute('data-bs-theme',t);try{localStorage.setItem('th',t)}catch(e){}})}
function setTheme(k){smooth(function(h){h.setAttribute('data-theme',k);try{localStorage.setItem('tm',k)}catch(e){}});drawDots()}
function drawDots(){var k=document.documentElement.getAttribute('data-theme')||'green';document.getElementById('dots').innerHTML=THEMES.map(function(t){return '<button type="button" class="dot'+(t.k==k?' on':'')+'" role="radio" aria-checked="'+(t.k==k)+'" aria-label="'+t.n+'" title="'+t.n+'" data-k="'+t.k+'" onclick="setTheme(this.dataset.k)" style="background:linear-gradient(135deg,'+t.c[0]+','+t.c[1]+')"></button>'}).join('');var c=THEMES.filter(function(t){return t.k==k})[0];document.getElementById('tname').textContent=c?c.n:''}
drawDots();
var CM=[['speedometer2','Dashboard','/'],['list-ul','Entries','/entries'],['table','Monthly Sheet','/sheet'],['tags','Edit Heads (rename / merge)','/heads'],['piggy-bank','Budget & Savings Goal','/budget'],['file-earmark-bar-graph','Monthly Report','/monthly'],['people','Loan Summary','/people'],['wallet2','Balance Summary (Cash/Bank/bKash/Loan)','/balance'],['download','Download Center (Excel / PDF)','/downloads'],['file-earmark-spreadsheet','All reports in one Excel file','#xa'],['journal-text','Journal','/journal'],['book','Ledger','/ledger'],['calculator','Trial Balance','/trial'],['plus-circle','Add new entry','#qa'],['file-earmark-excel','Download Excel (this page)','#xl'],['printer','Make PDF (Save as PDF)','#pr'],['cloud-arrow-down','Take backup now','/backup'],['circle-half','Dark / Light mode','#th']].concat(THEMES.map(function(t){return['palette2','Theme: '+t.n,'#t:'+t.k]}));
var pi=document.getElementById('pi'),pl=document.getElementById('pl');
function M(id){return bootstrap.Modal.getOrCreateInstance(document.getElementById(id))}
function pal(){M('pal').show()}
function rl(){var q=pi.value.toLowerCase();pl.innerHTML=CM.filter(function(c){return c[1].toLowerCase().indexOf(q)>-1}).map(function(c){return '<a href="#" class="list-group-item list-group-item-action border-0" data-c="'+c[2]+'"><i class="bi bi-'+c[0]+' me-2"></i>'+c[1]+'</a>'}).join('')}
function go(c){M('pal').hide();if(c=='#qa')setTimeout(function(){M('qa').show()},250);else if(c=='#pr')setTimeout(function(){print()},300);else if(c=='#th')tgl();else if(c.indexOf('#t:')==0)setTheme(c.slice(3));else if(c=='#xa')location='/export_all?m='+(new URLSearchParams(location.search).get('m')||'');else if(c=='#xl')location='/export?u='+encodeURIComponent(location.pathname+location.search);else location=c}
pl.onclick=function(e){var a=e.target.closest('a');if(a){e.preventDefault();go(a.dataset.c)}};
pi.oninput=rl;pi.onkeydown=function(e){if(e.key=='Enter'){var a=pl.querySelector('a');if(a)go(a.dataset.c)}};
document.getElementById('pal').addEventListener('shown.bs.modal',function(){pi.value='';rl();pi.focus()});
document.onkeydown=function(e){var t=e.target,tag=(t&&t.tagName)||'';if((e.ctrlKey||e.metaKey)&&e.key=='k'){e.preventDefault();pal();return}if(e.key=='n'&&!e.ctrlKey&&!e.metaKey&&!e.altKey&&!e.shiftKey&&!/INPUT|TEXTAREA|SELECT/.test(tag)&&!(t&&t.isContentEditable)&&!document.querySelector('.modal.show')){e.preventDefault();M('qa').show()}};
var q=document.getElementById('q');q.oninput=function(){var t=q.value.toLowerCase();document.querySelectorAll('main table tr').forEach(function(r){if(r.querySelector('td'))r.style.display=r.textContent.toLowerCase().indexOf(t)>-1?'':'none'})};
document.querySelectorAll('main .table:not(.sh) thead th').forEach(function(h){h.style.cursor='pointer';h.onclick=function(){var tb=h.closest('table'),all=[].slice.call(tb.querySelectorAll('tr')).filter(function(r){return r.querySelector('td')}),rs=all.filter(function(r){return !r.classList.contains('fw-bold')}),fx=all.filter(function(r){return r.classList.contains('fw-bold')}),ci=h.cellIndex,d=h.dataset.d=(h.dataset.d=='1'?'-1':'1');rs.sort(function(a,b){var x=a.cells[ci].textContent.trim(),y=b.cells[ci].textContent.trim(),nx=parseFloat(x.replace(/,/g,'')),ny=parseFloat(y.replace(/,/g,''));return(isNaN(nx)||isNaN(ny)?x.localeCompare(y):nx-ny)*d});rs.concat(fx).forEach(function(r){r.parentNode.appendChild(r)})}});
document.querySelectorAll('.stat .val').forEach(function(e){var v=parseFloat(e.textContent.replace(/,/g,'')),n=0;(function f(){n++;e.textContent=Math.round(v*Math.min(n/25,1)).toLocaleString('en-US');if(n<25)requestAnimationFrame(f)})()});
if(/[?&]print=1/.test(location.search)){var pt=((document.querySelector('.ttl')||{}).textContent||'Report').trim(),pm=new URLSearchParams(location.search).get('m')||'';document.title=(pt+' '+pm).trim();window.addEventListener('load',function(){setTimeout(function(){print()},800)})}
if(window.flatpickr)flatpickr('[data-date]',{altInput:true,altFormat:'d-m-Y',dateFormat:'Y-m-d',static:true,disableMobile:true});
var CATS=['খরচ','ধারে খরচ করলাম','আয়'],NAMES=['ধার নিলাম','ধারে খরচ করলাম','ধার পরিশোধ','ধার দিলাম','ধার ফেরত পেলাম'];
function syncS(){var t=document.querySelector('#fS [name=type]').value,c=CATS.indexOf(t)>-1,n=NAMES.indexOf(t)>-1,sc=document.getElementById('sCat'),sn=document.getElementById('sName');sc.classList.toggle('d-none',!c);sn.classList.toggle('d-none',!n);if(!c)sc.querySelector('input').value='';if(!n)sn.querySelector('input').value=''}
function syncRow(sel){var tr=sel.closest('tr'),t=sel.value,c=CATS.indexOf(t)>-1,n=NAMES.indexOf(t)>-1,nc=tr.querySelector('[name="category[]"]'),nn=tr.querySelector('[name="name[]"]');nc.readOnly=!c;nn.readOnly=!n;nc.style.opacity=c?1:.35;nn.style.opacity=n?1:.35;if(!c)nc.value='';if(!n)nn.value=''}
function qtab(x){var s=x=='s';document.getElementById('fS').classList.toggle('d-none',!s);document.getElementById('fB').classList.toggle('d-none',s);document.getElementById('tabB').classList.toggle('d-none',!s);document.getElementById('tabS').classList.toggle('d-none',s);document.getElementById('qtitle').textContent=s?'New Entry':'Bulk Entry';if(s)syncS();var pn=document.getElementById(s?'fS':'fB');pn.classList.remove('pane-in');void pn.offsetWidth;pn.classList.add('pane-in')}
function addRow(){var n=document.getElementById('rowtpl').content.firstElementChild.cloneNode(true);document.getElementById('rows').appendChild(n);syncRow(n.querySelector('[name="type[]"]'))}
for(var i=0;i<3;i++)addRow();
syncS();
document.getElementById('qa').addEventListener('hidden.bs.modal',function(){qtab('s')});
/* ---- v16: toast + ajax forms + styled confirm ---- */
var TI={warn:'exclamation-triangle-fill',success:'check-lg',error:'exclamation-triangle-fill',info:'info-lg',delete:'trash3-fill'};
function toast(msg,kind,ms){kind=TI[kind]?kind:'success';ms=ms||(kind=='error'?6500:3600);var box=document.getElementById('toasts');while(box.children.length>=4)box.firstChild.remove();var e=document.createElement('div');e.className='tst tst-'+kind;e.setAttribute('role',kind=='error'?'alert':'status');e.style.setProperty('--d',ms/1000+'s');e.innerHTML='<i class="bi bi-'+TI[kind]+'"></i><div class="tm"></div><button type="button" class="tcl" aria-label="বন্ধ করুন"><i class="bi bi-x-lg"></i></button><span class="tp"></span>';e.querySelector('.tm').textContent=msg;var shut=function(){if(e._s)return;e._s=1;e.classList.add('out');setTimeout(function(){e.remove()},220)};e.querySelector('.tcl').onclick=shut;e.querySelector('.tp').addEventListener('animationend',shut);box.appendChild(e)}
(window.FL||[]).forEach(function(x,i){setTimeout(function(){toast(x[1],x[0])},120+i*140)});
function sendForm(f,b){if(f._busy)return;f._busy=1;var h=b?b.innerHTML:'';if(b){b.style.minWidth=b.offsetWidth+'px';b.disabled=true;b.innerHTML='<span class="spinner-border spinner-border-sm" role="status"></span>'}var back=function(){f._busy=0;if(b){b.disabled=false;b.innerHTML=h}};fetch(f.getAttribute('action')||location.href,{method:'POST',body:new FormData(f),headers:{'X-Requested-With':'fetch'},credentials:'same-origin'}).then(function(r){return r.json().then(function(j){return{ok:r.ok&&j.ok,j:j}})}).then(function(x){if(x.ok){location.assign(x.j.next||'/')}else{toast(x.j.msg||'কিছু একটা ঠিক হয়নি, আবার চেষ্টা করুন','error');back()}}).catch(function(){toast('সার্ভারের সাথে যোগাযোগ হয়নি, আবার চেষ্টা করুন','error');back()})}
document.querySelectorAll('.pg[data-w]').forEach(function(e){setTimeout(function(){e.style.width=e.dataset.w+'%'},120)});
var cfPend=null;
document.addEventListener('submit',function(ev){var f=ev.target;if(!f.hasAttribute||!f.hasAttribute('data-ajax'))return;ev.preventDefault();var b=ev.submitter||f.querySelector('button:not([type=button])');if(f.hasAttribute('data-confirm')){var tr=f.closest('tr'),sm='';if(tr&&tr.cells.length>6)sm=tr.cells[0].textContent.trim()+', '+tr.cells[1].textContent.trim()+', '+tr.cells[6].textContent.trim()+' টাকা';document.getElementById('cft').textContent=f.getAttribute('data-confirm');document.getElementById('cfs').textContent=sm;cfPend=[f,b];M('cf').show();return}sendForm(f,b)});
document.getElementById('cfy').onclick=function(){var p=cfPend;cfPend=null;M('cf').hide();if(p)sendForm(p[0],p[1])};
document.getElementById('cf').addEventListener('shown.bs.modal',function(){document.getElementById('cfy').focus()});
document.getElementById('cf').addEventListener('hidden.bs.modal',function(){cfPend=null});
function stickyFix(){var f=document.querySelector('.fixbar'),m=document.querySelector('.main');if(!f||!m)return;m.style.paddingTop=getComputedStyle(f).position=='fixed'?(f.offsetHeight+16)+'px':''}
window.addEventListener('resize',stickyFix);window.addEventListener('load',stickyFix);stickyFix();
</script></body></html>"""

T = """{% macro tbl(es, d=False) %}<div class="table-responsive"><table class="table table-hover table-sm table-striped"><thead><tr><th>তারিখ<th>ধরন<th>মাধ্যম<th>খাত<th>বিবরণ<th>নাম<th class="text-end">টাকা{% if d %}<th>{% endif %}</tr></thead>
{% for e in es %}<tr><td>{{e.date|dmy}}<td><span class="badge text-bg-{{e.type|tc}}">{{e.type}}</span><td>{{e.medium}}<td>{{e.category or ''}}<td>{{e.description or ''}}<td>{{e.name or ''}}<td class="text-end">{{e.amount|n}}
{% if d %}<td class="text-nowrap"><a class="btn btn-sm btn-outline-primary" title="সম্পাদনা" href="/edit/{{e.id}}?next={{request.full_path|urlencode}}"><i class="bi bi-pencil"></i></a> <form class="d-inline" method="post" action="/del/{{e.id}}" data-ajax data-confirm="এই এন্ট্রি মুছবেন?"><button class="btn btn-sm btn-outline-danger" title="মুছুন"><i class="bi bi-trash"></i></button></form>{% endif %}</tr>{% endfor %}</table></div>{% endmacro %}
{% macro msel(ms, m, allow_all=True) %}{% set mm = m or curm() %}<div class="d-flex flex-wrap align-items-center gap-2 mb-3"><div class="mnav"><a class="mb-btn" href="{{ mq(shift(mm, -1)) }}" title="Previous month"><i class="bi bi-chevron-left"></i></a><span class="mb-label">{{ 'All time' if (allow_all and not m) else mm|mname }}</span><a class="mb-btn" href="{{ mq(shift(mm, 1)) }}" title="Next month"><i class="bi bi-chevron-right"></i></a></div>
{% if allow_all %}<a class="mb-chip {{'on' if not m}}" href="{{ mq('all') }}">All</a>{% endif %}<a class="mb-chip {{'on' if m == curm()}}" href="{{ mq(curm()) }}">This month</a></div>{% endmacro %}"""

def page(title, tmpl, **kw):
    pre, mark, post = tmpl.partition("<!--FIX-->")
    defs = "".join(re.findall(r"\{% macro .*?\{% endmacro %\}", pre, re.S)) if mark else ""
    fix = render_template_string(T + pre, **kw) if mark else ""
    body = render_template_string(T + defs + (post if mark else tmpl), **kw)
    return render_template_string(BASE, title=title, body=body, fix=fix)

@app.route("/")
def home():
    m = request.args.get("m"); m = curm() if m is None else ("" if m == "all" else m)
    ops, rows, bal = books(m or None)
    net = lambda a: bal[a][0] - bal[a][1]
    cash = {m: net(m) for m in MEDIA}
    pay = sum(-net(a) for a in list(bal) if kind(a) == "L")
    rec = sum(net(a) for a in list(bal) if a.startswith("Receivable"))
    cats = [r[0] for r in db().execute("SELECT DISTINCT category FROM entries WHERE category!='' ORDER BY 1")]
    names = [r[0] for r in db().execute("SELECT DISTINCT name FROM entries WHERE name!='' ORDER BY 1")]
    recent = db().execute("SELECT * FROM entries" + (" WHERE date LIKE ?" if m else "") + " ORDER BY id DESC LIMIT 10", (m + "%",) if m else ()).fetchall()
    mi, me, cm = defaultdict(float), defaultdict(float), defaultdict(float)
    for e in entries():
        k = e["date"][:7]
        if e["type"] == "আয়": mi[k] += e["amount"]
        if e["type"] in NEED_CAT: me[k] += e["amount"]
    last = m or max(list(mi) + list(me), default="")
    for e in entries(last or None):
        if e["type"] in NEED_CAT: cm[e["category"]] += e["amount"]
    top = sorted(cm.items(), key=lambda x: -x[1])[:8]; ks = sorted(set(mi) | set(me))[-12:]
    chart = dict(l=[k[5:7] + "-" + k[:4] for k in ks], i=[mi[k] for k in ks], e=[me[k] for k in ks], cl=[t[0] for t in top], cv=[t[1] for t in top], m=(last[5:7] + "-" + last[:4]) if last else "")
    mm = m or curm(); bud, _sp, _src = month_budget(mm); alerts = [b for b in bud if b["st"] != "ok"]
    gl = goal_get(mm)[0]; saved = mi.get(mm, 0) - me.get(mm, 0)
    hist = []
    for k in [shift(mm, -i) for i in range(5, -1, -1)]:
        v = mi.get(k, 0) - me.get(k, 0)
        gk = goal_get(k)[0]; hist.append(dict(k=k, v=v, cur=k == mm, l=calendar.month_abbr[int(k[5:7])], hp=max(0, min(100, v / gk * 100)) if gk else 0, full=bool(gk) and v >= gk, has=bool(gk)))
    sv = dict(g=gl, s=saved, p=max(0, saved / gl * 100) if gl else 0, inc=mi.get(mm, 0), exp=me.get(mm, 0), hist=hist)
    return page("Dashboard", """
{{ msel(ms, m) }}{% set icons = ['cash-stack','bank','phone','arrow-up-right-circle','arrow-down-left-circle','wallet2'] %}<div class="row g-3 mb-4">{% for l,v in cards %}<div class="col-6 col-md-4 col-xl-2"><div class="stat g{{loop.index}}"><i class="bi bi-{{icons[loop.index0]}}"></i><div class="lbl">{{l}}</div><div class="val">{{v|n}}</div></div></div>{% endfor %}</div><!--FIX--><div class="small text-muted mb-3">{{ 'নির্বাচিত মাস শেষ পর্যন্ত Cash/Bank/bKash ও Loan এর অবস্থা; নিচে শুধু ওই মাসের এন্ট্রি' if m else 'এখন পর্যন্ত সব এন্ট্রি অনুযায়ী অবস্থা (মাস বাছলে ওই মাস শেষের অবস্থা দেখাবে)' }}</div>
{% if alerts %}<div class="mb-3 d-flex flex-column gap-2">{% for a in alerts %}<div class="alert alert-{{'danger' if a.st=='over' else 'warning'}} d-flex align-items-center gap-2 py-2 mb-0"><i class="bi bi-{{'exclamation-octagon-fill' if a.st=='over' else 'exclamation-triangle-fill'}}"></i><div><b>{{a.c}}</b> — {% if a.st=='over' %}মাসিক বাজেট ছাড়িয়ে গেছে: {{a.s|n}} / {{a.l|n}} টাকা ({{(a.s-a.l)|n}} বেশি){% else %}বাজেটের {{a.p|round|int}}% খরচ হয়ে গেছে: {{a.s|n}} / {{a.l|n}} টাকা ({{(a.l-a.s)|n}} বাকি){% endif %}</div></div>{% endfor %}</div>{% endif %}
<div class="row g-3 mb-4"><div class="col-lg-5"><div class="card h-100"><div class="card-header d-flex align-items-center"><span><i class="bi bi-piggy-bank me-2"></i>সঞ্চয় লক্ষ্য — {{mm|mname}}</span><a class="ms-auto small" href="/budget?m={{mm}}">লক্ষ্য ঠিক করুন</a></div><div class="card-body">
{% if sv.g %}<div class="d-flex align-items-baseline gap-2 flex-wrap"><span class="sv-big">{{sv.s|n}}</span><span class="text-muted">/ {{sv.g|n}} টাকা</span><span class="ms-auto badge text-bg-{{'success' if sv.p>=100 else 'secondary'}}">{{sv.p|round|int}}%</span></div>
<div class="pgt my-3"><div class="pg {{'ok done' if sv.p>=100 else 'ok'}}" data-w="{{[sv.p,100]|min}}"></div></div>
<div class="small {{'text-success fw-semibold' if sv.p>=100 else 'text-muted'}}">{% if sv.s<0 %}এই মাসে খরচ আয়ের চেয়ে {{(-sv.s)|n}} টাকা বেশি, তাই এখনো কিছু জমেনি।{% elif sv.p>=100 %}লক্ষ্য পূর্ণ হয়েছে, {{(sv.s-sv.g)|n}} টাকা বেশি জমেছে।{% else %}আর {{(sv.g-sv.s)|n}} টাকা জমালেই লক্ষ্য পূর্ণ।{% endif %}</div>
<div class="small text-muted">সঞ্চয় = আয় {{sv.inc|n}} − খরচ {{sv.exp|n}}</div>
<div class="sv-hist mt-3">{% for h in sv.hist %}<div class="sv-col{{' cur' if h.cur}}" title="{{h.k|mname}}: {{h.v|n}} টাকা{{' (লক্ষ্য ঠিক করা ছিল না)' if not h.has}}"><div class="sv-bar"><i class="{{'full' if h.full else 'part'}}" style="height:{{h.hp}}%"></i></div><span>{{h.l}}</span></div>{% endfor %}</div>
{% else %}<div class="text-center text-muted py-4"><i class="bi bi-piggy-bank fs-1 d-block mb-2"></i>মাসে কত টাকা জমাতে চান তা ঠিক করলে এখানে প্রতি মাসের অগ্রগতি দেখাবে।<div class="mt-3"><a class="btn btn-sm btn-primary" href="/budget?m={{mm}}">সঞ্চয় লক্ষ্য ঠিক করুন</a></div></div>{% endif %}</div></div></div>
<div class="col-lg-7"><div class="card h-100"><div class="card-header d-flex align-items-center"><span><i class="bi bi-speedometer me-2"></i>খাতভিত্তিক বাজেট — {{mm|mname}}</span><a class="ms-auto small" href="/budget?m={{mm}}">বাজেট ঠিক করুন</a></div><div class="card-body">
{% for b in bud[:6] %}<div class="bd-row"><div class="d-flex align-items-baseline gap-2"><span class="fw-semibold">{{b.c}}</span><span class="small text-muted">{{b.s|n}} / {{b.l|n}}</span><span class="ms-auto small fw-semibold bd-{{b.st}}">{{b.p|round|int}}%{% if b.st=='over' %} · {{(b.s-b.l)|n}} বেশি{% endif %}</span></div><div class="pgt sm"><div class="pg {{b.st}}" data-w="{{[b.p,100]|min}}"></div></div></div>{% else %}<div class="text-center text-muted py-4"><i class="bi bi-speedometer fs-1 d-block mb-2"></i>কোনো খাতের বাজেট ঠিক করা নেই। বাজেট ঠিক করলে সীমা পার হওয়ার আগেই সতর্কতা পাবেন।<div class="mt-3"><a class="btn btn-sm btn-primary" href="/budget?m={{mm}}">বাজেট ঠিক করুন</a></div></div>{% endfor %}
{% if bud|length>6 %}<div class="small text-muted mt-2">আরও {{bud|length-6}}টি খাত — <a href="/budget?m={{mm}}">সব দেখুন</a></div>{% endif %}</div></div></div></div>
<div class="row g-3 mb-4"><div class="col-lg-8"><div class="card h-100"><div class="card-header">মাসভিত্তিক আয় বনাম খরচ</div><div class="card-body"><canvas id="c1" height="130"></canvas></div></div></div>
<div class="col-lg-4"><div class="card h-100"><div class="card-header">খাতভিত্তিক খরচ ({{chart.m}})</div><div class="card-body"><canvas id="c2"></canvas></div></div></div></div>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.4/dist/chart.umd.min.js"></script><script>
var C={{chart|tojson}};function cv(n){return getComputedStyle(document.documentElement).getPropertyValue(n).trim()}
function cpal(){return[cv('--b1'),cv('--b2'),'#f59e0b','#0ea5e9','#e11d48','#8b5cf6','#14b8a6','#64748b']}
Chart.defaults.color=cv('--muted');Chart.defaults.font.family="'Inter','Hind Siliguri',sans-serif";
var ch1=new Chart(document.getElementById('c1'),{type:'bar',data:{labels:C.l,datasets:[{label:'আয়',data:C.i,backgroundColor:cv('--b2'),borderRadius:6},{label:'খরচ',data:C.e,backgroundColor:'#f59e0b',borderRadius:6}]},options:{plugins:{legend:{position:'bottom'}},scales:{y:{grid:{color:'rgba(148,163,184,.2)'}}}}});
var ch2=new Chart(document.getElementById('c2'),{type:'doughnut',data:{labels:C.cl,datasets:[{data:C.cv,backgroundColor:cpal(),borderColor:cv('--card')}]},options:{plugins:{legend:{position:'bottom'}}}});
window.reTheme=function(){Chart.defaults.color=cv('--muted');ch1.data.datasets[0].backgroundColor=cv('--b2');ch2.data.datasets[0].backgroundColor=cpal();ch2.data.datasets[0].borderColor=cv('--card');ch1.update();ch2.update()};
window.addEventListener('load',function(){var A={{alerts|tojson}},o=A.filter(function(a){return a.st=='over'});if(!o.length)return;var k='bo'+{{mm|tojson}}+o.length;try{if(sessionStorage.getItem(k))return;sessionStorage.setItem(k,1)}catch(e){}toast(o.length==1?o[0].c+' খাতে মাসিক বাজেট ছাড়িয়ে গেছে':o.length+' টি খাতে মাসিক বাজেট ছাড়িয়ে গেছে','warn',6500)});
</script>
<h6 class="fw-bold mb-3">সাম্প্রতিক ১০টি এন্ট্রি</h6>{{ tbl(recent, True) }}""",
        cards=[("Cash", cash["Cash"]), ("Bank", cash["Bank"]), ("bKash", cash["bKash"]), ("Payable (দেব)", pay), ("Receivable (পাব)", rec), ("Net", sum(cash.values()) + rec - pay)],
        today=date.today().isoformat(), types=TYPES, media=MEDIA, cats=cats, names=names, recent=recent, chart=chart, ms=months(), m=m, mm=mm, bud=bud, alerts=alerts, sv=sv)

@app.route("/add", methods=["POST"])
def add():
    f = request.form; t = f["type"]
    cat, nm = f.get("category", "").strip(), f.get("name", "").strip().title()
    try: amt = float(f["amount"])
    except ValueError: amt = 0
    if amt <= 0 or (t in NEED_CAT and not cat) or (t in NEED_NAME and not nm):
        if ajax(): return jsonify(ok=False, msg="টাকা দিন, আর এই ধরনের জন্য খাত/নাম আবশ্যক।"), 400
        return page("ভুল", "<div class='alert alert-danger'>টাকা, এবং এই ধরনের জন্য খাত/নাম দিতে হবে। <a href='/'>ফিরে যান</a></div>")
    db().execute("INSERT INTO entries(date,type,medium,category,description,name,amount) VALUES(?,?,?,?,?,?,?)",
                 (f["date"], t, f["medium"], cat, f.get("description", "").strip(), nm, amt))
    db().commit(); r = done("এন্ট্রি সংরক্ষণ হয়েছে", safe_next(f.get("next", "/")))
    for w in budget_warn([(f["date"], t, cat, amt)]): flash(w, "warn")
    return r

@app.route("/add_bulk", methods=["POST"])
def add_bulk():
    f = request.form; d = f.get("date", "").strip()
    if re.fullmatch(r"\d{2}-\d{2}-\d{4}", d): d = d[6:] + "-" + d[3:5] + "-" + d[:2]
    if not d: d = date.today().isoformat()
    cols = ["type", "medium", "category", "name", "description", "amount"]
    L = {k: f.getlist(k + "[]") for k in cols}
    rows, bad = [], []
    for i in range(len(L["amount"])):
        g = {k: (L[k][i] if i < len(L[k]) else "").strip() for k in cols}
        if not any(g[k] for k in ("amount", "category", "name", "description")): continue
        try: amt = float(g["amount"])
        except ValueError: amt = 0
        t, cat, nm = g["type"], g["category"], g["name"].title()
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", d) or t not in TYPES or g["medium"] not in MEDIA or amt <= 0 or (t in NEED_CAT and not cat) or (t in NEED_NAME and not nm):
            bad.append(str(i + 1)); continue
        rows.append((d, t, g["medium"], cat, g["description"], nm, amt))
    if bad:
        if ajax(): return jsonify(ok=False, msg="সারি %s এ সমস্যা আছে — টাকা এবং ধরন অনুযায়ী খাত/নাম দেখুন। কিছুই সংরক্ষণ হয়নি।" % ", ".join(bad)), 400
        return page("Error", "<div class='alert alert-danger'>Problem in row(s): %s<br>Check Amount, and Category/Name for the selected Type. Nothing was saved. <a href='javascript:history.back()'>Go back to fix</a></div>" % ", ".join(bad))
    if rows:
        db().executemany("INSERT INTO entries(date,type,medium,category,description,name,amount) VALUES(?,?,?,?,?,?,?)", rows); db().commit()
        r = done("%d টি এন্ট্রি সংরক্ষণ হয়েছে" % len(rows), safe_next(f.get("next", "/")))
        for w in budget_warn([(x[0], x[1], x[3], x[6]) for x in rows]): flash(w, "warn")
        return r
    return done("সংরক্ষণ করার মতো কোনো সারি পাওয়া যায়নি", safe_next(f.get("next", "/")), "info")

@app.route("/del/<int:i>", methods=["POST"])
def delete(i):
    db().execute("DELETE FROM entries WHERE id=?", (i,)); db().commit(); return done("এন্ট্রি মুছে ফেলা হয়েছে", request.referrer or "/entries", "delete")

@app.route("/entries")
def all_entries():
    m = getm()
    return page("Entries", "{{ msel(ms, m) }}<!--FIX-->{{ tbl(es, True) }}", ms=months(), m=m, es=entries(m or None))

@app.route("/journal")
def journal():
    m = getm()
    rows = [r for r in books()[1] if r[0]["date"].startswith(m)]
    return page("Journal", """{{ msel(ms, m) }}<!--FIX--><div class="table-responsive"><table class="table table-hover table-sm table-striped"><thead><tr><th>তারিখ<th>Dr (ডেবিট)<th>Cr (ক্রেডিট)<th class="text-end">টাকা<th>বিবরণ</tr></thead>
{% for e,d,c in rows %}<tr><td>{{e.date|dmy}}<td>{{d}}<td>{{c}}<td class="text-end">{{e.amount|n}}<td>{{e.type}} {{e.description or ''}}</tr>{% endfor %}</table></div>""", ms=months(), m=m, rows=rows)

@app.route("/ledger")
def ledger():
    m = getm()
    ops, rows, bal = books()
    accs = sorted(bal, key=lambda a: (kind(a), a)); acc = request.args.get("acc", accs[0] if accs else "")
    sign = 1 if kind(acc) in "AX" else -1
    run = sign * ops.get(acc, 0.0); jer = run; lines = []
    if not m and run: lines.append(("", "প্রারম্ভিক ব্যালেন্স", ops[acc], 0, run))
    for e, d, c in rows:
        if acc not in (d, c): continue
        dr, cr = (e["amount"], 0) if acc == d else (0, e["amount"]); k = e["date"][:7]
        if m and k > m: continue
        run += sign * (dr - cr)
        if m and k < m: jer = run; continue
        lines.append((e["date"], f'{e["type"]} {e["description"] or ""} {e["name"] or ""}', dr, cr, run))
    if m: lines.insert(0, ("", "আগের জের (Opening)", 0, 0, jer))
    return page("Ledger", """<form class="mb-2 d-flex flex-wrap gap-2"><select name="acc" class="form-select w-auto" onchange="this.form.submit()">{% for a in accs %}<option {{'selected' if a==acc}}>{{a}}</option>{% endfor %}</select>
<select name="m" class="form-select w-auto" onchange="this.form.submit()"><option value="">সব মাস</option>{% for x in ms %}<option value="{{x}}" {{'selected' if x==m}}>{{x|my}}</option>{% endfor %}</select></form>
<table class="table table-hover table-sm"><thead><tr><th>তারিখ<th>বিবরণ<th class="text-end">Dr<th class="text-end">Cr<th class="text-end">ব্যালেন্স</tr></thead>
{% for d,t,dr,cr,b in lines %}<tr><td>{{d|dmy}}<td>{{t}}<td class="text-end">{{dr|n if dr else ''}}<td class="text-end">{{cr|n if cr else ''}}<td class="text-end">{{b|n}}</tr>{% endfor %}</table>""", accs=accs, acc=acc, lines=lines, ms=months(), m=m)

@app.route("/trial")
def trial():
    m = getm(); bal = books(m or None)[2]; out = []; td = tc = 0
    for a in sorted(bal, key=lambda a: ("ALEIX".index(kind(a)), a)):
        n = bal[a][0] - bal[a][1]
        if abs(n) < 0.005: continue
        dr, cr = (n, 0) if n > 0 else (0, -n); td += dr; tc += cr; out.append((a, KN[kind(a)], dr, cr))
    return page("Trial Balance", """{{ msel(ms, m) }}<!--FIX--><div class="small text-muted mb-2">{{ "নির্বাচিত মাস শেষ পর্যন্ত" if m else "সব এন্ট্রি অনুযায়ী (এখন পর্যন্ত)" }}</div><table class="table table-hover table-sm table-striped"><thead><tr><th>অ্যাকাউন্ট<th>ধরন<th class="text-end">Dr<th class="text-end">Cr</tr></thead>
{% for a,k,d,c in out %}<tr><td>{{a}}<td>{{k}}<td class="text-end">{{d|n if d else ''}}<td class="text-end">{{c|n if c else ''}}</tr>{% endfor %}
<tr class="fw-bold"><td colspan="2">মোট<td class="text-end">{{td|n}}<td class="text-end">{{tc|n}}</tr></table>
<div class="alert {{'alert-success' if ok else 'alert-danger'}}">{{'✔ ট্রায়াল ব্যালেন্স মিলেছে' if ok else '✘ মিলছে না — কোথাও ভুল আছে'}}</div>""", out=out, td=td, tc=tc, ms=months(), m=m, ok=abs(td - tc) < 0.005)

HEADS = ["মাছ", "মুরগি", "সবজি", "ফলমূল", "ডিম", "চাল", "মুদি বাজার (তেল/লবণ/মশলা)", "দুধ", "পান", "শুকনা বাজার", "খাওয়া দাওয়া (বাহিরে)", "গ্যাস", "পানি (বিল/জার)", "মোবাইল রিচার্জ", "মেডিসিন", "ডাক্তার", "যাতায়াত (রিকশা/সিএনজি/বাস)", "বিকাশ খরচ", "কিতাব/শিক্ষা", "লন্ড্রি", "কুরবানী", "কসাই", "আটা", "দান", "ব্যক্তিগত খরচ", "অন্যান্য", "টেস্ট", "হাবু ভাই", "মধু", "মেহমানদারী", "আতর", "আরমান"]

@app.route("/people")
@app.route("/loan")
@app.route("/loans")
@app.route("/loan-summary")
def people():
    m = getm()
    p = defaultdict(lambda: defaultdict(float)); o = defaultdict(lambda: [0.0, 0.0])
    for e in entries():
        n, k, t, a = e["name"], e["date"][:7], e["type"], e["amount"]
        if not n or (m and k > m): continue
        if m and k < m:
            if t in ("ধার নিলাম", "ধারে খরচ করলাম"): o[n][0] += a
            elif t == "ধার পরিশোধ": o[n][0] -= a
            elif t == "ধার দিলাম": o[n][1] += a
            elif t == "ধার ফেরত পেলাম": o[n][1] -= a
        else: p[n][t] += a
    out = []
    for n in dict.fromkeys(list(o) + list(p)):
        v = p[n]; b, x, r, l, f = v["ধার নিলাম"], v["ধারে খরচ করলাম"], v["ধার পরিশোধ"], v["ধার দিলাম"], v["ধার ফেরত পেলাম"]
        op, orc = o[n]; py, rc = op + b + x - r, orc + l - f
        row = (n, op, b, x, r, py, orc, l, f, rc, rc - py)
        if any(abs(z) > .005 for z in row[1:]): out.append(row)
    bal = books(m or None)[2]; net = lambda a: bal[a][0] - bal[a][1]
    cash = sum(net(md) for md in MEDIA)
    pay, rec = sum(z[5] for z in out), sum(z[9] for z in out)
    ok = abs(pay - sum(-net(a) for a in list(bal) if kind(a) == "L")) < .5 and abs(rec - sum(net(a) for a in list(bal) if a.startswith("Receivable"))) < .5
    return page("Loan Summary", """{{ msel(ms, m) }}<!--FIX--><div class="small text-muted mb-3">{{ 'নির্বাচিত মাসের লেনদেন, আগের মাসের জের (Opening) ও মাস শেষের বাকি' if m else 'সব মাস মিলিয়ে এখন পর্যন্ত (মাস বাছলে ওই মাসের হিসাব দেখাবে)' }}</div>
<div class="row g-2 mb-3">{% for l,v in cards %}<div class="col-6 col-md"><div class="card text-center"><div class="card-body p-2"><div class="small text-muted">{{l}}</div><div class="fs-5 fw-bold">{{v|n}}</div></div></div></div>{% endfor %}
<div class="col-6 col-md"><div class="card text-center"><div class="card-body p-2"><div class="small text-muted">যাচাই</div><div class="fs-6 fw-bold">{{'✔ সব মিলেছে' if ok else '✘ মিলছে না'}}</div></div></div></div></div>
<div class="small text-muted mb-1">নামের উপর ক্লিক করলে ওই ব্যক্তির পুরো লেনদেন (লেজার) দেখা যাবে। Payable ও Receivable Cash এর সাথে মেশানো হয়নি।</div>
<div class="table-responsive"><table class="table table-hover table-sm"><thead><tr><th>#<th>নাম{% if m %}<th class="text-end">Opening Payable{% endif %}<th class="text-end">ধার নিয়েছি<th class="text-end">ধারে খরচ করলাম<th class="text-end">পরিশোধ করেছি<th class="text-end">বাকি Payable (দেব){% if m %}<th class="text-end">Opening Receivable{% endif %}<th class="text-end">ধার দিয়েছি<th class="text-end">ফেরত পেয়েছি<th class="text-end">বাকি Receivable (পাব)<th class="text-end">Net (পাওনা − দেনা)</tr></thead>
{% for r in out %}<tr><td>{{loop.index}}<td><a href="/ledger?acc={{ (('Payable: ' if r[2]+r[3]+r[1] else 'Receivable: ') ~ r[0])|urlencode }}">{{r[0]}}</a>{% if m %}<td class="text-end">{{r[1]|n}}{% endif %}<td class="text-end">{{r[2]|n}}<td class="text-end">{{r[3]|n}}<td class="text-end">{{r[4]|n}}<td class="text-end fw-bold">{{r[5]|n}}{% if m %}<td class="text-end">{{r[6]|n}}{% endif %}<td class="text-end">{{r[7]|n}}<td class="text-end">{{r[8]|n}}<td class="text-end fw-bold">{{r[9]|n}}<td class="text-end fw-bold">{{r[10]|n}}</tr>{% endfor %}
<tr class="fw-bold"><td><td>মোট{% for i in range(1,11) %}{% if m or i not in (1,6) %}<td class="text-end">{{ out|sum(attribute=i)|n }}{% endif %}{% endfor %}</tr></table></div>""",
        out=out, ok=ok, m=m, ms=months(), cards=[("মোট Payable (আমি দেব)", pay), ("মোট Receivable (আমি পাব)", rec), ("Cash (হাতে)", cash), ("Net (Cash + পাব − দেব)", cash + rec - pay)])

DEFAULT_CATS = ["মাছ", "মুরগি", "সবজি", "ফলমূল", "ডিম", "চাল", "মুদি বাজার (তেল/লবণ/মশলা)", "দুধ", "পান", "শুকনা বাজার", "খাওয়া দাওয়া (বাহিরে)", "গ্যাস", "পানি (বিল/জার)", "মোবাইল রিচার্জ", "মেডিসিন", "ডাক্তার", "যাতায়াত (রিকশা/সিএনজি/বাস)", "বিকাশ খরচ", "কিতাব/শিক্ষা", "লন্ড্রি", "কুরবানী", "কসাই", "আটা", "দান", "ব্যক্তিগত খরচ", "অন্যান্য", "টেস্ট", "হাবু ভাই", "মধু", "মেহমানদারী", "আতর", "আরমান"]

@app.route("/sheet")
@app.route("/monthsheet")
@app.route("/monthlysheet")
@app.route("/month-sheet")
@app.route("/monthly-sheet")
@app.route("/monthly_sheet")
def sheet():
    import calendar
    ms = sorted(set(months() + [date.today().strftime("%Y-%m")]), reverse=True)
    m = getm() or ms[0]; hide = request.args.get("hide")
    allE = entries(); ops = sum(r["amount"] for r in db().execute("SELECT * FROM opening"))
    opening = ops + sum(CASH[e["type"]] * e["amount"] for e in allE if e["date"][:7] < m)
    es = [e for e in allE if e["date"][:7] == m]
    ren = dict((r[0], r[1]) for r in db().execute("SELECT old, new FROM renames")); base = list(dict.fromkeys(ren.get(c, c) for c in DEFAULT_CATS))
    cats = base + [c for c in dict.fromkeys(e["category"] for e in es if e["type"] in NEED_CAT) if c not in base]
    y, mm = map(int, m.split("-")); day = {d: defaultdict(float) for d in range(1, calendar.monthrange(y, mm)[1] + 1)}
    for e in es:
        v = day[int(e["date"][8:10])]
        if e["type"] in NEED_CAT: v["c:" + e["category"]] += e["amount"]; v["exp"] += e["amount"]
        v[e["type"]] += e["amount"]; v["cash"] += CASH[e["type"]] * e["amount"]
    tot = defaultdict(float)
    for v in day.values():
        for k, x in v.items(): tot[k] += x
    if hide: cats = [c for c in cats if tot["c:" + c]]
    rows, cur = [], opening
    for d, v in day.items():
        inh = cur; cur += v["cash"]
        rows.append(("%s-%02d" % (m, d), [v["c:" + c] for c in cats], v["ধার পরিশোধ"], v["ধার দিলাম"], v["ধার নিলাম"], v["আয়"], v["exp"], inh, cur, v["ধার ফেরত পেলাম"], v["ধারে খরচ করলাম"]))
    totrow = ("Total", [tot["c:" + c] for c in cats], tot["ধার পরিশোধ"], tot["ধার দিলাম"], tot["ধার নিলাম"], tot["আয়"], tot["exp"], None, cur, tot["ধার ফেরত পেলাম"], tot["ধারে খরচ করলাম"])
    inc, ex, cr = tot["আয়"], tot["exp"], tot["ধারে খরচ করলাম"]
    cin, cout = inc + tot["ধার নিলাম"] + tot["ধার ফেরত পেলাম"], ex - cr + tot["ধার পরিশোধ"] + tot["ধার দিলাম"]
    upto = [e for e in allE if e["date"][:7] <= m]; S = lambda t: sum(e["amount"] for e in upto if e["type"] == t)
    top = max(((c, tot["c:" + c]) for c in cats), key=lambda x: x[1], default=("", 0))
    summ = [("আয় ও ব্যয়", None), ("মোট আয় (Received)", inc), ("মোট খরচ (Total Expense)", ex), ("   এর মধ্যে ধারে কেনা (এখনও পরিশোধ হয়নি)", cr), ("   এর মধ্যে নগদে পরিশোধিত", ex - cr), ("আয় − খরচ (উদ্বৃত্ত / ঘাটতি)", inc - ex),
            ("ধার ও অন্যান্য লেনদেন", None), ("ধার নেওয়া (Loan Taken)", tot["ধার নিলাম"]), ("ধার পরিশোধ করেছি (Payable Paid)", tot["ধার পরিশোধ"]), ("ধার দিয়েছি (Loan Given)", tot["ধার দিলাম"]), ("ধার ফেরত পেয়েছি (Loan Recovered)", tot["ধার ফেরত পেলাম"]),
            ("Cash (শুধু প্রকৃত হাতে থাকা টাকা)", None), ("মাসের শুরুর Cash", opening), ("  + মোট নগদ জমা (আয় + ধার নেওয়া + ধার ফেরত)", cin), ("  − মোট নগদ বের (নগদ খরচ + ধার পরিশোধ + ধার দেওয়া)", cout), ("মাস শেষে Cash (Remaining)", cur),
            ("যাচাই", "✔ মিলেছে" if abs(opening + cin - cout - cur) < .5 else "✘ মিলছে না"),
            ("মাস শেষে বকেয়া (Cash-এর সাথে যুক্ত নয়)", None), ("মোট Payable (আমি দেব)", S("ধার নিলাম") + S("ধারে খরচ করলাম") - S("ধার পরিশোধ")), ("মোট Receivable (আমি পাব)", S("ধার দিলাম") - S("ধার ফেরত পেলাম")),
            ("অন্যান্য", None), ("এই মাসের এন্ট্রি সংখ্যা", len(es)), ("সর্বোচ্চ খরচের খাত", top[0]), ("সর্বোচ্চ খাতের খরচ", top[1])]
    return page("Monthly Sheet", """<style>:root{--bd:#c5cedd}[data-bs-theme=dark]{--bd:#425372}
.sw{max-height:calc(100vh - 210px);overflow:auto;border-radius:14px;box-shadow:var(--shadow)}
.table.sh{border-collapse:collapse;border-radius:0;box-shadow:none;border:0;margin:0}
.table.sh td,.table.sh th{white-space:nowrap;font-size:.82rem;border:1px solid var(--bd);padding:.32rem .6rem;line-height:1.25}
.table.sh thead th{position:sticky;top:0;z-index:2;background:#dcecc9;color:#334155;text-align:center;vertical-align:middle}
[data-bs-theme=dark] .table.sh thead th{background:#1f3a26;color:#e2e8f0}
.table.sh td:first-child,.table.sh th:first-child{position:sticky;left:0;z-index:1;background:var(--card);font-weight:600;border-right:2px solid var(--bd)}
.table.sh thead th:first-child{z-index:3}
.table.sh tbody tr:nth-child(even) td:first-child{background:#f0f6ea}
[data-bs-theme=dark] .table.sh tbody tr:nth-child(even) td:first-child{background:#183021}
.table.sh tr.fw-bold td{background:rgba(44,87,42,.12)!important;border-top:2px solid var(--bd)}@page{size:A4 landscape;margin:6mm}</style>
<div class="d-flex flex-wrap align-items-center gap-3">{{ msel(ms, m, False) }}<form class="mb-3"><input type="hidden" name="m" value="{{m}}"><label class="small"><input type="checkbox" name="hide" value="1" {{'checked' if hide}} onchange="this.form.submit()"> ফাঁকা খাত লুকান</label></form></div>
<!--FIX--><div class="sw"><table class="table table-hover table-sm table-bordered sh"><thead><tr><th>{{m|my}}{% for c in cats %}<th>{{c}}{% endfor %}<th>Payable Paid<th>Loan Given<th>Loan Taken<th>Received<th>Total Expense<th>Cash In Hand<th>Remaining<th>Loan Recovered<th>Credit Purchase (বাকি)</tr></thead>
<tr class="fw-bold"><td>Opening Balance{% for c in cats %}<td>{% endfor %}<td><td><td><td><td><td><td class="text-end">{{opening|n}}<td><td></tr>
{% for r in rows + [tot] %}<tr {{'class=fw-bold' if loop.last}}><td>{{r[0]|dmy}}{% for v in r[1] %}<td class="text-end">{{v|z}}{% endfor %}{% for v in r[2:] %}<td class="text-end">{% if v is none %}{% elif loop.index in (6,7) %}{{v|n}}{% else %}{{v|z}}{% endif %}{% endfor %}</tr>{% endfor %}</table></div>""",
        ms=ms, m=m, hide=hide, cats=cats, rows=rows, tot=totrow, opening=opening, summ=summ)

BK = os.path.join(os.path.dirname(DB), "backups")
def do_backup():
    os.makedirs(BK, exist_ok=True)
    p = os.path.join(BK, "hisab-%s.db" % date.today().isoformat())
    src, dst = sqlite3.connect(DB), sqlite3.connect(p); src.backup(dst); dst.close(); src.close()
    for old in sorted(os.listdir(BK))[:-30]: os.remove(os.path.join(BK, old))
    return p
_last = [None]
@app.before_request
def auto_backup():
    if _last[0] != date.today():
        _last[0] = date.today()
        try: do_backup()
        except Exception: pass

@app.route("/backup")
def backup():
    return send_file(do_backup(), as_attachment=True, download_name="hisab-backup-%s.db" % date.today().isoformat())

@app.context_processor
def inject():
    d = db()
    return dict(qc=[r[0] for r in d.execute("SELECT DISTINCT category FROM entries WHERE category!='' ORDER BY 1")],
                qn=[r[0] for r in d.execute("SELECT DISTINCT name FROM entries WHERE name!='' ORDER BY 1")],
                today=date.today().isoformat(), types=TYPES, media=MEDIA, fname=os.path.basename(__file__))

class TP(HTMLParser):
    def __init__(s):
        super().__init__(); s.tables = []; s.row = None; s.cell = None; s.skip = 0; s.ttl = ""; s.inttl = False
    def fc(s):
        if s.cell is not None and s.row is not None: s.row.append(" ".join("".join(s.cell).split()))
        s.cell = None
    def fr(s):
        s.fc()
        if s.row and s.tables: s.tables[-1].append(s.row)
        s.row = None
    def handle_starttag(s, t, a):
        if t == "table": s.tables.append([])
        elif t == "tr" and s.tables: s.fr(); s.row = []
        elif t in ("td", "th") and s.row is not None: s.fc(); s.cell = []
        elif t in ("form", "style", "script"): s.skip += 1
        elif t == "div" and "ttl" in (dict(a).get("class") or ""): s.inttl = True
    def handle_endtag(s, t):
        if t in ("tr", "table"): s.fr()
        elif t in ("form", "style", "script"): s.skip = max(0, s.skip - 1)
        elif t == "div": s.inttl = False
    def handle_data(s, d):
        if s.inttl: s.ttl += d
        elif s.cell is not None and not s.skip: s.cell.append(d)

def num(x):
    t = x.replace(",", "")
    return float(t) if t and t.lstrip("-").replace(".", "", 1).isdigit() else x

def have_openpyxl():
    try:
        import openpyxl; return True
    except ImportError:
        try:
            import subprocess, sys, importlib
            subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
            importlib.invalidate_caches(); import openpyxl; return True
        except Exception:
            return False

def fill_sheet(ws, u):
    from openpyxl.styles import Font, PatternFill, Border, Side
    p = TP(); p.feed(app.test_client().get(u).get_data(as_text=True))
    ws.append(["হিসাব — " + p.ttl.strip(), date.today().strftime("%d-%m-%Y")]); ws["A1"].font = Font(bold=True, size=14); ws.append([])
    th = Side(style="thin", color="BBBBBB")
    for tb in p.tables:
        for ri, row in enumerate(tb):
            while row and row[-1] == "": row.pop()
            ws.append([num(x) for x in row])
            for c in ws[ws.max_row]:
                c.border = Border(left=th, right=th, top=th, bottom=th)
                if ri == 0: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E4E6FA")
                if isinstance(c.value, float): c.number_format = "#,##0"
        ws.append([])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(38, max(10, max((len(str(c.value or "")) for c in col[2:]), default=0) + 2))

def send_wb(wb, name):
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=name, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

NOXL = "<div class='alert alert-warning'>Excel ফাইল বানানোর টুল (openpyxl) নিজে থেকে ইনস্টল করা গেল না। ইন্টারনেট চালু আছে কিনা দেখে আবার চেষ্টা করুন। না হলে Command Prompt এ <b>python -m pip install openpyxl</b> লিখে Enter চাপুন।</div>"

@app.route("/export")
def export():
    if not have_openpyxl(): return page("Excel", NOXL)
    from openpyxl import Workbook
    u = request.args.get("u", "/")
    if not u.startswith("/") or u.startswith(("/export", "/backup")): return redirect("/")
    wb = Workbook(); wb.active.title = "Hisab"; fill_sheet(wb.active, u)
    return send_wb(wb, "hisab-%s-%s.xlsx" % (re.sub(r"[^A-Za-z0-9]+", "-", u).strip("-") or "home", date.today().isoformat()))

@app.route("/export_all")
def export_all():
    if not have_openpyxl(): return page("Excel", NOXL)
    from openpyxl import Workbook
    m = getm(); m = m if re.fullmatch(r"\d{4}-\d{2}", m) else ""
    q = "m=%s&" % m if m else ""
    wb = Workbook(); wb.remove(wb.active)
    for nm, u in [("Monthly Sheet", "/sheet?%shide=1" % q), ("Monthly Report", "/monthly?%s" % q), ("Loan Summary", "/people" + ("?m=" + m if m else "")), ("Balance Summary", "/balance" + ("?m=" + m if m else ""))]:
        fill_sheet(wb.create_sheet(nm), u)
    return send_wb(wb, "hisab-all-reports-%s.xlsx" % (m or date.today().isoformat()))

@app.route("/monthly")
@app.route("/report")
@app.route("/monthlyreport")
@app.route("/monthly-report")
def report():
    ms = sorted(set(months() + [date.today().strftime("%Y-%m")]), reverse=True)
    m = getm() or ms[0]
    allE = entries(); ops = sum(r["amount"] for r in db().execute("SELECT * FROM opening"))
    opening = ops + sum(CASH[e["type"]] * e["amount"] for e in allE if e["date"][:7] < m)
    exp, inc, tk, pd, t = defaultdict(float), defaultdict(float), defaultdict(lambda: [0.0, 0.0]), defaultdict(float), defaultdict(float)
    for e in allE:
        if e["date"][:7] != m: continue
        a, ty, n = e["amount"], e["type"], e["name"]; t[ty] += a
        if ty in NEED_CAT: exp[e["category"] or "খাত নেই"] += a
        if ty == "আয়": inc[e["category"] or "উৎস দেওয়া হয়নি"] += a
        if ty == "ধার নিলাম": tk[n][0] += a
        if ty == "ধারে খরচ করলাম": tk[n][1] += a
        if ty == "ধার পরিশোধ": pd[n] += a
    od = defaultdict(float)
    for e in allE:
        if e["date"][:7] < m:
            if e["type"] in ("ধার নিলাম", "ধারে খরচ করলাম"): od[e["name"]] += e["amount"]
            elif e["type"] == "ধার পরিশোধ": od[e["name"]] -= e["amount"]
    names = [n for n in dict.fromkeys(list(od) + list(tk) + list(pd)) if n]
    taken = sorted([(n, tk[n][0], tk[n][1], sum(tk[n])) for n in names if sum(tk[n])], key=lambda r: -r[3])
    paid = sorted([(n, pd[n]) for n in names if pd[n]], key=lambda r: -r[1])
    due = [(n, od[n], sum(tk[n]), pd[n], od[n] + sum(tk[n]) - pd[n]) for n in names if od[n] or sum(tk[n]) or pd[n]]
    ie, ee = sorted(inc.items(), key=lambda x: -x[1]), sorted(exp.items(), key=lambda x: -x[1])
    ti, te, cr = t["আয়"], sum(exp.values()), t["ধারে খরচ করলাম"]
    lt, lp, lg, lr = t["ধার নিলাম"], t["ধার পরিশোধ"], t["ধার দিলাম"], t["ধার ফেরত পেলাম"]
    cin, cout = opening + ti + lt + lr, (te - cr) + lp + lg; rem = cin - cout
    rec = sum(e["amount"] * (1 if e["type"] == "ধার দিলাম" else -1) for e in allE if e["date"][:7] <= m and e["type"] in ("ধার দিলাম", "ধার ফেরত পেলাম"))
    taken_total, due_total = sum(r[3] for r in taken), sum(r[4] for r in due)
    mo, allm, c2, d2 = defaultdict(lambda: defaultdict(float)), [], ops, 0.0
    for e in allE: mo[e["date"][:7]][e["type"]] += e["amount"]
    for k in sorted(mo):
        v = mo[k]; c2 += sum(CASH[x] * v[x] for x in CASH); tk2 = v["ধার নিলাম"] + v["ধারে খরচ করলাম"]; d2 += tk2 - v["ধার পরিশোধ"]
        allm.append((k, v["আয়"], v["খরচ"] + v["ধারে খরচ করলাম"], tk2, v["ধার পরিশোধ"], c2, d2))
    st = [("Opening Balance (আগের মাস থেকে)", opening, "s"), ("+ Total Income", ti, "i"), ("+ Loan Taken (নগদ ধার)", lt, "i"), ("+ Loan Recovered (ধার ফেরত পেয়েছি)", lr, "i"),
          ("= মোট নগদ জমা", cin, "s"), ("− নগদ খরচ (Total Expenses − বাকিতে কেনা)", te - cr, "o"), ("− Loan Paid (ধার পরিশোধ)", lp, "o"), ("− Loan Given (ধার দিয়েছি)", lg, "o"),
          ("= মোট নগদ বের", cout, "s"), ("Remaining (মাস শেষে নগদ)", rem, "t")]
    st = [r for r in st if r[1] or r[2] in "st"]
    cards = [("Opening Balance", opening, "calendar2-check", 2), ("Total Income", ti, "graph-up-arrow", 5), ("Total Expenses", te, "graph-down-arrow", 4), ("আয় − খরচ", ti - te, "arrow-left-right", 6), ("Remaining (নগদ)", rem, "wallet2", 1)]
    cards2 = [("Total Loan Taken", taken_total, "box-arrow-in-down", 4), ("Total Loan Paid", lp, "box-arrow-up", 2), ("Total Due Remaining", due_total, "hourglass-split", 3)]
    return page("Monthly Report", """
{% macro head(title, rows, tot, icon) %}<div class="card h-100"><div class="card-header"><i class="bi bi-{{icon}} me-2"></i>{{title}}</div><div class="card-body p-0"><table class="table table-sm mb-0 shadow-none border-0"><thead><tr><th>#<th>হেড<th class="text-end">টাকা<th class="text-end">%</tr></thead>
{% for c,v in rows %}<tr><td>{{loop.index}}<td>{{c}}<td class="text-end">{{v|n}}<td class="text-end">{{((v/tot*100)|round(1)) if tot else 0}}%</tr>{% endfor %}
<tr class="fw-bold"><td><td>মোট<td class="text-end">{{tot|n}}<td></tr></table></div></div>{% endmacro %}
{{ msel(ms, m, False) }}<!--FIX-->
{% for grp in (cards, cards2) %}<div class="row g-3 mb-3">{% for l,v,ic,g in grp %}<div class="col-6 col-lg"><div class="stat g{{g}}"><i class="bi bi-{{ic}}"></i><div class="lbl">{{l}}</div><div class="val">{{v|n}}</div></div></div>{% endfor %}</div>{% endfor %}
<div class="row g-3 mb-3"><div class="col-lg-6">{{ head('Head Wise Expense (খাতভিত্তিক খরচ)', ee, te, 'graph-down-arrow') }}</div><div class="col-lg-6">{{ head('Head Wise Income (উৎসভিত্তিক আয়)', ie, ti, 'graph-up-arrow') }}</div></div>
<div class="row g-3 mb-3"><div class="col-lg-6"><div class="card h-100"><div class="card-header"><i class="bi bi-cash-coin me-2"></i>নগদ হিসাব (Income, Expenses, Remaining)</div><div class="card-body p-0"><table class="table table-sm mb-0 shadow-none border-0"><thead><tr><th>বিবরণ<th class="text-end">টাকা</tr></thead>
{% for l,v,k in st %}<tr class="{{'fw-bold' if k in 'st'}}"><td>{{l}}<td class="text-end {{'text-success' if k=='i' else 'text-danger' if k=='o' else ''}}">{{v|n}}</tr>{% endfor %}
<tr><td>Total Income − Total Expenses (বাকি সহ)<td class="text-end">{{(ti-te)|n}}</tr></table></div></div></div>
<div class="col-lg-6"><div class="card h-100"><div class="card-header"><i class="bi bi-arrow-right-circle me-2"></i>পরের মাসে স্বয়ংক্রিয়ভাবে যাবে</div><div class="card-body"><div class="small text-muted">Opening Balance (নগদ)</div><div class="fs-3 fw-bold">{{rem|n}}</div>
<div class="small text-muted mt-3">Opening Due (দেনা)</div><div class="fs-4 fw-bold">{{due_total|n}}</div>{% if rec %}<div class="small text-muted mt-3">পাওনা (Receivable)</div><div class="fs-5 fw-bold">{{rec|n}}</div>{% endif %}
<div class="small text-muted mt-3">এই মাসের Remaining ও Due নিজে থেকেই পরের মাসের Opening হয়ে যায়, আলাদা এন্ট্রি দিতে হয় না। আগের কোনো মাসের এন্ট্রি বদলালে পরের সব মাসের Opening ঠিক হয়ে যাবে।</div></div></div></div></div>
<div class="row g-3 mb-3"><div class="col-lg-7"><div class="card h-100"><div class="card-header"><i class="bi bi-box-arrow-in-down me-2"></i>Head Wise Loan Taken (যার কাছ থেকে নিয়েছি)</div><div class="card-body p-0"><table class="table table-sm mb-0 shadow-none border-0"><thead><tr><th>#<th>নাম<th class="text-end">নগদ ধার<th class="text-end">ধারে খরচ (বাকি)<th class="text-end">মোট</tr></thead>
{% for n,a,b,c in taken %}<tr><td>{{loop.index}}<td>{{n}}<td class="text-end">{{a|z}}<td class="text-end">{{b|z}}<td class="text-end">{{c|n}}</tr>{% endfor %}
<tr class="fw-bold"><td><td>Total Loan Taken<td class="text-end">{{lt|n}}<td class="text-end">{{cr|n}}<td class="text-end">{{taken_total|n}}</tr></table></div></div></div>
<div class="col-lg-5"><div class="card h-100"><div class="card-header"><i class="bi bi-box-arrow-up me-2"></i>Head Wise Loan Paid (পরিশোধ)</div><div class="card-body p-0"><table class="table table-sm mb-0 shadow-none border-0"><thead><tr><th>#<th>নাম<th class="text-end">পরিশোধ</tr></thead>
{% for n,a in paid %}<tr><td>{{loop.index}}<td>{{n}}<td class="text-end">{{a|n}}</tr>{% endfor %}
<tr class="fw-bold"><td><td>Total Loan Paid<td class="text-end">{{lp|n}}</tr></table></div></div></div></div>
<div class="card mb-3"><div class="card-header"><i class="bi bi-hourglass-split me-2"></i>Total Due Remaining (নাম অনুযায়ী বাকি)</div><div class="card-body p-0"><table class="table table-sm mb-0 shadow-none border-0"><thead><tr><th>#<th>নাম<th class="text-end">Opening Due<th class="text-end">+ এই মাসে Taken<th class="text-end">− এই মাসে Paid<th class="text-end">Closing Due</tr></thead>
{% for n,o,a,b,c in due %}<tr><td>{{loop.index}}<td>{{n}}<td class="text-end">{{o|z}}<td class="text-end">{{a|z}}<td class="text-end">{{b|z}}<td class="text-end fw-bold">{{c|n}}</tr>{% endfor %}
<tr class="fw-bold"><td><td>মোট<td class="text-end">{{due|sum(attribute=1)|n}}<td class="text-end">{{taken_total|n}}<td class="text-end">{{lp|n}}<td class="text-end">{{due_total|n}}</tr></table></div></div>
<div class="card mb-3"><div class="card-header"><i class="bi bi-calendar3 me-2"></i>সব মাসের সারসংক্ষেপ (মাসে ক্লিক করলে সেই মাসের রিপোর্ট)</div><div class="card-body p-0"><table class="table table-sm mb-0 shadow-none border-0"><thead><tr><th>মাস<th class="text-end">Total Income<th class="text-end">Total Expenses<th class="text-end">Loan Taken<th class="text-end">Loan Paid<th class="text-end">মাস শেষে Remaining (নগদ)<th class="text-end">মাস শেষে Due</tr></thead>
{% for r in allm %}<tr><td><a href="/monthly?m={{r[0]}}">{{r[0]}}</a>{% for v in r[1:] %}<td class="text-end">{{v|n}}{% endfor %}</tr>{% endfor %}
<tr class="fw-bold"><td>মোট<td class="text-end">{{allm|sum(attribute=1)|n}}<td class="text-end">{{allm|sum(attribute=2)|n}}<td class="text-end">{{allm|sum(attribute=3)|n}}<td class="text-end">{{allm|sum(attribute=4)|n}}<td><td></tr></table></div></div>""",
        ms=ms, m=m, ee=ee, ie=ie, te=te, ti=ti, st=st, rem=rem, due_total=due_total, rec=rec, taken=taken, paid=paid, due=due, lt=lt, lp=lp, cr=cr, taken_total=taken_total, cards=cards, cards2=cards2, allm=allm)

@app.route("/edit/<int:i>", methods=["GET", "POST"])
def edit(i):
    e = db().execute("SELECT * FROM entries WHERE id=?", (i,)).fetchone()
    nx = request.values.get("next") or "/entries"
    if not nx.startswith("/") or nx.startswith("//"): nx = "/entries"
    if not e: return redirect(nx)
    err = ""
    if request.method == "POST":
        f = request.form; t = f.get("type"); cat = f.get("category", "").strip(); nm = f.get("name", "").strip().title()
        try: amt = float(f.get("amount", ""))
        except ValueError: amt = 0
        if t not in TYPES or f.get("medium") not in MEDIA or amt <= 0 or (t in NEED_CAT and not cat) or (t in NEED_NAME and not nm):
            err = "টাকা দিতে হবে, আর এই ধরনের জন্য খাত/নাম আবশ্যক।"
            if ajax(): return jsonify(ok=False, msg=err), 400
            e = dict(id=i, date=f.get("date"), type=t, medium=f.get("medium"), category=cat, description=f.get("description", ""), name=nm, amount=f.get("amount"))
        else:
            db().execute("UPDATE entries SET date=?, type=?, medium=?, category=?, description=?, name=?, amount=? WHERE id=?",
                         (f["date"], t, f["medium"], cat, f.get("description", "").strip(), nm, amt, i)); db().commit()
            return done("এন্ট্রি #%d আপডেট হয়েছে" % i, nx)
    return page("Edit Entry", """{% if err %}<div class="alert alert-danger">{{err}}</div>{% endif %}
<div class="card" style="max-width:820px"><div class="card-header"><i class="bi bi-pencil-square me-2"></i>এন্ট্রি #{{e.id}} ঠিক করুন</div><div class="card-body"><form method="post" data-ajax class="row g-3"><input type="hidden" name="next" value="{{nx}}">
<div class="col-md-4"><label class="small text-muted">তারিখ</label><input type="text" data-date name="date" value="{{e.date}}" class="form-control" required></div>
<div class="col-md-4"><label class="small text-muted">ধরন</label><select name="type" class="form-select">{% for t in types %}<option {{'selected' if t==e.type}}>{{t}}</option>{% endfor %}</select></div>
<div class="col-md-4"><label class="small text-muted">মাধ্যম</label><select name="medium" class="form-select">{% for m in media %}<option {{'selected' if m==e.medium}}>{{m}}</option>{% endfor %}</select></div>
<div class="col-md-4"><label class="small text-muted">খাত / আয়ের উৎস</label><input name="category" list="qc" value="{{e.category or ''}}" class="form-control"></div>
<div class="col-md-4"><label class="small text-muted">নাম (ধার/বাকির ক্ষেত্রে)</label><input name="name" list="qn" value="{{e.name or ''}}" class="form-control"></div>
<div class="col-md-4"><label class="small text-muted">টাকা</label><input type="number" step="any" name="amount" value="{{e.amount}}" class="form-control" required></div>
<div class="col-12"><label class="small text-muted">বিবরণ</label><input name="description" value="{{e.description or ''}}" class="form-control"></div>
<div class="col-12 d-flex gap-2"><button class="btn btn-primary">সংরক্ষণ করুন</button><a class="btn btn-outline-secondary" href="{{nx}}">বাতিল</a></div></form>
<div class="small text-muted mt-3">সংরক্ষণ করলে জার্নাল, লেজার, ট্রায়াল ব্যালেন্স, Monthly Sheet ও Monthly Report সব নিজে থেকেই ঠিক হয়ে যাবে। পরের মাসের Opening-ও আপডেট হবে।</div></div></div>""", e=e, nx=nx, err=err)

@app.route("/balance")
def balance():
    d = request.args.get("d", ""); m = getm()
    if not re.fullmatch(r"\d{4}-\d{2}", m): m = ""
    if m and not d: d = "%s-%02d" % (m, calendar.monthrange(int(m[:4]), int(m[5:7]))[1])
    cash = {r["medium"]: r["amount"] for r in db().execute("SELECT * FROM opening")}
    pay, rec = defaultdict(float), defaultdict(float)
    for e in entries():
        if d and e["date"] > d: continue
        a, t = e["amount"], e["type"]; cash[e["medium"]] = cash.get(e["medium"], 0) + CASH[t] * a
        if t in ("ধার নিলাম", "ধারে খরচ করলাম"): pay[e["name"]] += a
        elif t == "ধার পরিশোধ": pay[e["name"]] -= a
        elif t == "ধার দিলাম": rec[e["name"]] += a
        elif t == "ধার ফেরত পেলাম": rec[e["name"]] -= a
    rows = [(n, pay[n], rec[n], rec[n] - pay[n]) for n in dict.fromkeys(list(pay) + list(rec)) if n and (abs(pay[n]) > .005 or abs(rec[n]) > .005)]
    hand, tp, tr = sum(cash.values()), sum(r[1] for r in rows), sum(r[2] for r in rows)
    pos = [("Cash", cash.get("Cash", 0)), ("Bank", cash.get("Bank", 0)), ("bKash", cash.get("bKash", 0)), ("মোট হাতে আছে (Cash + Bank + bKash)", hand),
           ("Total Loan Payable (আমি দেব)", tp), ("Total Loan Receivable (আমি পাব)", tr), ("Net (হাতে আছে + পাব − দেব)", hand + tr - tp)]
    cards = [(p[0], p[1], ic, g) for p, ic, g in zip([pos[0], pos[1], pos[2], pos[3], pos[4], pos[5], pos[6]], ["cash-stack", "bank", "phone", "wallet2", "arrow-up-right-circle", "arrow-down-left-circle", "calculator"], [1, 2, 3, 6, 4, 5, 6])]
    return page("Balance Summary", """<form class="mb-3 d-flex flex-wrap gap-2 align-items-center"><span class="small text-muted">অবস্থা:</span><select name="m" class="form-select w-auto" onchange="this.form.d.value='';this.form.submit()"><option value="">মাস বাছুন</option>{% for x in ms %}<option value="{{x}}" {{'selected' if x==m}}>{{x|my}}</option>{% endfor %}</select><span class="small text-muted">অথবা তারিখ:</span><input type="text" data-date name="d" value="{{d}}" class="form-control w-auto" onchange="this.form.m.value='';this.form.submit()">
<span class="small text-muted">{{ (d|dmy) + ' তারিখ পর্যন্ত' if d else 'এখন পর্যন্ত সব এন্ট্রি অনুযায়ী' }}</span>{% if d %}<a href="/balance">সব দেখুন</a>{% endif %}</form>
<div class="row g-3 mb-4">{% for l,v,ic,g in cards %}<div class="col-6 col-md-4 col-xl"><div class="stat g{{g}}"><i class="bi bi-{{ic}}"></i><div class="lbl">{{l}}</div><div class="val">{{v|n}}</div></div></div>{% endfor %}</div>
<div class="row g-3"><div class="col-lg-5"><div class="card h-100"><div class="card-header"><i class="bi bi-wallet2 me-2"></i>এখন কোথায় কত আছে</div><div class="card-body p-0"><table class="table table-sm mb-0 shadow-none border-0"><thead><tr><th>বিবরণ<th class="text-end">টাকা</tr></thead>
{% for l,v in pos %}<tr class="{{'fw-bold' if loop.index in (4,7)}}"><td>{{l}}<td class="text-end">{{v|n}}</tr>{% endfor %}</table></div></div></div>
<div class="col-lg-7"><div class="card h-100"><div class="card-header"><i class="bi bi-people me-2"></i>Loan — নাম অনুযায়ী</div><div class="card-body p-0"><table class="table table-sm mb-0 shadow-none border-0"><thead><tr><th>#<th>নাম<th class="text-end">Payable (আমি দেব)<th class="text-end">Receivable (আমি পাব)<th class="text-end">Net (পাব − দেব)</tr></thead>
{% for n,p,r,x in rows %}<tr><td>{{loop.index}}<td>{{n}}<td class="text-end">{{p|z}}<td class="text-end">{{r|z}}<td class="text-end">{{x|n}}</tr>{% endfor %}
<tr class="fw-bold"><td><td>মোট<td class="text-end">{{tp|n}}<td class="text-end">{{tr|n}}<td class="text-end">{{(tr-tp)|n}}</tr></table></div></div></div></div>""", d=d, m=m, ms=months(), cards=cards, pos=pos, rows=rows, tp=tp, tr=tr)

@app.route("/downloads")
def downloads():
    ms = sorted(set(months() + [date.today().strftime("%Y-%m")]), reverse=True)
    m = getm() or ms[0]
    items = [("Monthly Sheet", "table", "প্রতিদিনের খাতভিত্তিক শিট (ফাঁকা খাত বাদে)", "/sheet?m=%s&hide=1" % m, True),
             ("Monthly Report", "file-earmark-bar-graph", "Head wise Income/Expense, Loan, Due ও Remaining", "/monthly?m=%s" % m, True),
             ("Loan Summary", "people", "নাম অনুযায়ী Payable ও Receivable (মাস শেষ পর্যন্ত)", "/people?m=%s" % m, True),
             ("Balance Summary", "wallet2", "Cash / Bank / bKash ও Loan (মাস শেষ পর্যন্ত)", "/balance?m=%s" % m, True)]
    return page("Download Center", """{{ msel(ms, m, False) }}<!--FIX--><div class="small text-muted mb-3">Monthly Sheet, Monthly Report, Loan Summary ও Balance Summary — বাছা মাস অনুযায়ী</div>
<div class="row g-3 mb-3">{% for t,ic,desc,u,mo in items %}<div class="col-md-6"><div class="card h-100"><div class="card-body"><div class="d-flex align-items-center gap-3 mb-3"><div class="fs-2" style="color:var(--ac)"><i class="bi bi-{{ic}}"></i></div><div><div class="fw-bold fs-5">{{t}}{% if mo %} <span class="badge text-bg-secondary">{{m|my}}</span>{% endif %}</div><div class="small text-muted">{{desc}}</div></div></div>
<a class="btn btn-success" href="/export?u={{u|urlencode}}"><i class="bi bi-file-earmark-excel"></i> Excel ডাউনলোড</a> <a class="btn btn-danger" target="_blank" href="{{u}}{{'&' if '?' in u else '?'}}print=1"><i class="bi bi-file-earmark-pdf"></i> PDF</a></div></div></div>{% endfor %}</div>
<div class="card"><div class="card-body d-flex flex-wrap align-items-center gap-3"><div class="me-auto"><div class="fw-bold">৪টি রিপোর্ট একসাথে (এক Excel ফাইলে ৪টি শিট)</div><div class="small text-muted">Monthly Sheet, Monthly Report, Loan Summary, Balance Summary</div></div><a class="btn btn-primary" href="/export_all?m={{m}}"><i class="bi bi-download"></i> সব একসাথে Excel</a></div></div>
<div class="small text-muted mt-3">PDF বাটনে ক্লিক করলে নতুন ট্যাবে প্রিন্ট উইন্ডো নিজে থেকে খুলবে। সেখানে <b>Destination / Printer</b> এ <b>Save as PDF</b> বেছে <b>Save</b> চাপুন। ফাইলের নাম আগে থেকেই বসানো থাকবে।</div>""", ms=ms, m=m, items=items)

@app.route("/budget", methods=["GET", "POST"])
def budget():
    if request.method == "POST":
        f = request.form; mo = f.get("month", "")
        if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", mo): mo = curm()
        db().execute("DELETE FROM mbudgets WHERE month=?", (mo,)); db().execute("INSERT INTO mbudgets VALUES(?,'',0)", (mo,))  # '' = এই মাস নিজে সেট করা
        for c_, a_ in zip(f.getlist("cat[]"), f.getlist("amt[]")):
            c_ = c_.strip()
            if not c_: continue
            try: a = float(a_)
            except ValueError: a = 0
            if a > 0: db().execute("INSERT OR REPLACE INTO mbudgets VALUES(?,?,?)", (mo, c_, a))
        try: gv = float(f.get("goal", "").strip() or 0)
        except ValueError: gv = 0
        db().execute("INSERT OR REPLACE INTO mgoals VALUES(?,?)", (mo, max(gv, 0)))  # ০ = এই মাসে লক্ষ্য নেই
        db().commit()
        return done("%s মাসের বাজেট ও সঞ্চয় লক্ষ্য সংরক্ষণ হয়েছে" % my(mo), safe_next(f.get("next"), "/budget"))
    m = getm()
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", m): m = curm()
    bud, spent, src = month_budget(m); have = {b["c"] for b in bud}
    rows = bud + [dict(c=c_, l=0, s=sp, p=0, st="none") for c_, sp in sorted(spent.items(), key=lambda x: -x[1]) if c_ and c_ not in have]
    gl, gsrc = goal_get(m); saved = sum(e["amount"] * (1 if e["type"] == "আয়" else -1) for e in entries(m) if e["type"] == "আয়" or e["type"] in NEED_CAT)
    ms = sorted(set(months() + [curm()]), reverse=True)
    return page("Budget & Goal", """{{ msel(ms, m, False) }}<!--FIX--><div class="small text-muted mb-3">প্রতিটি মাসের বাজেট আলাদা — উপরে মাস বদলে সেই মাসের বাজেট ঠিক করুন। সংরক্ষণ করলে শুধু <b>{{m|my}}</b> মাস বদলায়, অন্য মাস নয়। বাজেট ফাঁকা বা ০ রাখলে ওই খাতের বাজেট থাকে না। সঞ্চয় লক্ষ্যও প্রতি মাসে আলাদা।</div>
{% if src and src!=m %}<div class="alert alert-info py-2 small"><i class="bi bi-info-circle me-1"></i>{{m|my}} মাসের জন্য আলাদা বাজেট এখনো সেট করা নেই। {{ 'আগের সাধারণ বাজেট' if src=='0000-00' else (src|my ~ ' মাসের বাজেট') }} দেখানো হচ্ছে। সংরক্ষণ করলে এই মাসের নিজস্ব বাজেট তৈরি হবে।</div>{% endif %}
<form method="post" data-ajax><input type="hidden" name="next" value="{{request.full_path}}"><input type="hidden" name="month" value="{{m}}">
<div class="row g-3"><div class="col-lg-4"><div class="card h-100"><div class="card-header"><i class="bi bi-piggy-bank me-2"></i>মাসিক সঞ্চয় লক্ষ্য</div><div class="card-body">
<label class="small text-muted mb-1">{{m|my}} মাসে কত টাকা জমাতে চান</label><input type="number" step="any" min="0" name="goal" value="{{ (gl|int if gl==gl|int else gl) if gl else '' }}" class="form-control" placeholder="যেমন 10000">
{% if gsrc and gsrc!=m and gl %}<div class="small text-muted mt-2"><i class="bi bi-info-circle me-1"></i>এই মাসের নিজস্ব লক্ষ্য নেই; {{ 'আগের সাধারণ লক্ষ্য' if gsrc=='0000-00' else (gsrc|my ~ ' মাসের লক্ষ্য') }} দেখানো হচ্ছে। সংরক্ষণ করলে এই মাসের নিজস্ব হবে।</div>{% endif %}
<div class="small text-muted mt-3">{{m|mname}} মাসে সঞ্চয় (আয় − খরচ): <b class="{{'text-danger' if saved<0}}">{{saved|n}}</b> টাকা</div>
{% if gl %}<div class="pgt my-3"><div class="pg ok{{' done' if saved>=gl}}" data-w="{{[[saved/gl*100,0]|max,100]|min}}"></div></div><div class="small text-muted">লক্ষ্যের {{[saved/gl*100,0]|max|round|int}}% পূর্ণ</div>{% endif %}</div></div></div>
<div class="col-lg-8"><div class="card h-100"><div class="card-header"><i class="bi bi-speedometer me-2"></i>খরচের খাতের মাসিক বাজেট</div><div class="card-body p-0"><div class="table-responsive"><table class="table table-sm mb-0 shadow-none border-0 align-middle"><thead><tr><th>খাত<th class="text-end">{{m|my}} এ খরচ<th style="width:170px">মাসিক বাজেট (টাকা)<th style="min-width:170px">অবস্থা</tr></thead>
{% for r in rows %}<tr><td>{{r.c}}<input type="hidden" name="cat[]" value="{{r.c}}"><td class="text-end">{{r.s|n}}<td><input type="number" step="any" min="0" name="amt[]" value="{{r.l|int if r.l else ''}}" class="form-control form-control-sm text-end" placeholder="বাজেট নেই">
<td>{% if r.st!='none' %}<div class="pgt sm"><div class="pg {{r.st}}" data-w="{{[r.p,100]|min}}"></div></div><div class="small bd-{{r.st}}">{{r.p|round|int}}%{% if r.st=='over' %} · {{(r.s-r.l)|n}} বেশি{% elif r.st=='warn' %} · সীমার কাছে{% endif %}</div>{% else %}<span class="small text-muted">—</span>{% endif %}</tr>{% endfor %}
{% for i in range(2) %}<tr><td><input name="cat[]" list="qc" class="form-control form-control-sm" placeholder="নতুন খাতের নাম"><td class="text-end text-muted">—<td><input type="number" step="any" min="0" name="amt[]" class="form-control form-control-sm text-end" placeholder="বাজেট"><td></tr>{% endfor %}</table></div></div></div></div></div>
<div class="mt-3"><button class="btn btn-primary"><i class="bi bi-check2-circle me-1"></i>সংরক্ষণ করুন</button></div></form>
<div class="small text-muted mt-3">নতুন খরচ সেভ করার সময় কোনো খাতের বাজেটের ৮০% ছুঁলে বা ছাড়ালে সাথে সাথে টোস্ট সতর্কতা দেখাবে, আর ড্যাশবোর্ডেও ব্যানার থাকবে।</div>""", ms=ms, m=m, rows=rows, gl=gl, saved=saved, src=src, gsrc=gsrc)

@app.route("/heads", methods=["GET", "POST"])
def heads():
    if request.method == "POST":
        f = request.form; kd, old, new = f["kind"], f["old"], f["new"].strip(); ty = tuple(NEED_CAT) if kd == "exp" else ("আয়",)
        qs = ",".join("?" * len(ty)); n = mg = 0
        if new and new != old:
            mg = db().execute("SELECT COUNT(*) FROM entries WHERE category=? AND type IN (%s)" % qs, (new,) + ty).fetchone()[0]
            n = db().execute("UPDATE entries SET category=? WHERE category=? AND type IN (%s)" % qs, (new, old) + ty).rowcount
            if kd == "exp":
                for mo_, a_ in db().execute("SELECT month, amount FROM mbudgets WHERE category=?", (old,)).fetchall():
                    nb = db().execute("SELECT amount FROM mbudgets WHERE month=? AND category=?", (mo_, new)).fetchone()
                    db().execute("DELETE FROM mbudgets WHERE month=? AND category=?", (mo_, old)); db().execute("INSERT OR REPLACE INTO mbudgets VALUES(?,?,?)", (mo_, new, a_ + (nb[0] if nb else 0)))
                db().execute("UPDATE renames SET new=? WHERE new=?", (new, old)); db().execute("INSERT OR REPLACE INTO renames VALUES(?,?)", (old, new))
            db().commit()
        if n: flash("%d টি এন্ট্রির খাতের নাম বদলানো হয়েছে%s" % (n, " (আগে থেকে থাকা খাতের সাথে মিলিয়ে এক করা হয়েছে)" if mg else ""), "success")
        else: flash("কোনো পরিবর্তন হয়নি", "info")
        return redirect("/heads")
    def lst(ty):
        return db().execute("SELECT category, COUNT(*) c, SUM(amount) s FROM entries WHERE type IN (%s) AND category!='' GROUP BY category ORDER BY s DESC" % ",".join("?" * len(ty)), ty).fetchall()
    n = request.args.get("n"); msg = ""
    if n is not None: msg = ("✔ %s টি এন্ট্রির খাতের নাম বদলানো হয়েছে" % n) + (" (আগে থেকে থাকা খাতের সাথে মিলিয়ে এক করা হয়েছে)" if request.args.get("mg") == "1" else "") if n != "0" else "কোনো পরিবর্তন হয়নি"
    return page("Edit Heads", """{% macro hl(title, rows, kd, icon) %}<div class="card mb-3"><div class="card-header"><i class="bi bi-{{icon}} me-2"></i>{{title}}</div><div class="card-body p-0"><table class="table table-sm mb-0 shadow-none border-0"><thead><tr><th>#<th>বর্তমান নাম<th class="text-end">এন্ট্রি<th class="text-end">মোট টাকা<th>নতুন নাম লিখে বদলান</tr></thead>
{% for r in rows %}<tr><td>{{loop.index}}<td>{{r.category}}<td class="text-end">{{r.c}}<td class="text-end">{{r.s|n}}<td><form method="post" class="d-flex gap-2"><input type="hidden" name="kind" value="{{kd}}"><input type="hidden" name="old" value="{{r.category}}"><input name="new" value="{{r.category}}" class="form-control form-control-sm" required><button class="btn btn-sm btn-primary text-nowrap">বদলান</button></form></tr>{% endfor %}</table></div></div>{% endmacro %}
{% if msg %}<div class="alert alert-success">{{msg}}</div>{% endif %}
<div class="small text-muted mb-3">খাতের নাম বদলালে ওই খাতের সব এন্ট্রি, Monthly Sheet এর কলাম, Report ও চার্ট নিজে থেকেই বদলে যায়। আগে থেকে থাকা অন্য খাতের নাম লিখলে দুটো খাত এক হয়ে যায় (merge)। নতুন খাত যোগ করতে চাইলে এন্ট্রি দেওয়ার সময় খাতের ঘরে নতুন নাম লিখলেই হবে।</div>
{{ hl('খরচের খাত (Head Wise Expense)', exp, 'exp', 'graph-down-arrow') }}{{ hl('আয়ের উৎস (Head Wise Income)', inc, 'inc', 'graph-up-arrow') }}""", exp=lst(tuple(NEED_CAT)), inc=lst(("আয়",)), msg=msg)

@app.errorhandler(404)
def not_found(_):
    p = request.path.lower().replace(" ", "").replace("%20", "")
    if "sheet" in p: return redirect("/sheet")
    if "loan" in p: return redirect("/people")
    return page("পেজ নেই", "<div class='alert alert-warning'>এই ঠিকানায় কোনো পেজ নেই। উপরের মেনু থেকে পেজ বেছে নিন, অথবা <a href='/sheet'>মাসিক শিট</a> / <a href='/people'>লোন সামারি</a> এ যান।</div>"), 404

if __name__ == "__main__":
    threading.Timer(1.2, lambda: webbrowser.open("http://localhost:5000")).start()
    app.run(host="0.0.0.0", port=5000)
